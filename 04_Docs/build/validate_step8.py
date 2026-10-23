"""
Step 8 validation - documentation, deck and GitHub readiness.
  A. Deliverables   : README, LICENSE, .gitignore, user guide (MD + PDF, max 2 pages), dashboard preview PDF (8 pages),
                      deck (PPTX with 5 slides + PDF), 8 page screenshots + star schema
  B. Links          : every link / image in README and user guide resolves; every path in backticks and in the
                      repository tree exists; external links are on an allow-list of known sources
  C. Number audit   : every number in README, user guide and deck (incl. speaker notes) is in the Step 8 manifest or a
                      justified literal; every manifest entry re-checked against an INDEPENDENT source:
                        facts      -> insights_facts.json value + half-up formatting
                        measures   -> the value Power BI itself returned (step5_live_results.json), in the model's format
                        project    -> recounted differently (page folders, visual files, TMDL lines, pandas on the CSV)
                        validation -> PASS lines counted in each report
  D. DAX reference  : README's measure table and DAX code == the measures as Power BI reports them (name, format,
                      description, expression) - all of them, in order
  E. Deck           : 5 slides, a title on each, speaker notes on each, chart values == facts, image embedded,
                      no text outside the slide
  F. Consistency    : every "N/N checks" claim and every page / visual / measure / query count in ALL Markdown docs
                      of the project matches the reports and the project files
  G. Git readiness  : what `git add` would commit (dry run, separate git dir): no Power BI cache / local settings,
                      no leftovers, no file > 50 MB, no secrets or personal data, README explains ProjectFolder
Usage: python 04_Docs/build/validate_step8.py <project_root> <report.txt>
"""
import glob, json, os, re, shutil, subprocess, sys, tempfile, zipfile
from decimal import Decimal, ROUND_HALF_UP
import xml.etree.ElementTree as ET
import pandas as pd

ROOT, REPORT = sys.argv[1], sys.argv[2]
res, out = [], []
def check(sec, name, ok, detail=""):
    res.append(bool(ok)); out.append(f"[{'PASS' if ok else 'FAIL'}] {sec}: {name}" + (f" - {detail}" if detail else ""))
def P(*a): return os.path.join(ROOT, *a)
def read(rel): return open(P(rel), encoding="utf-8-sig").read()
def pdf_pages(rel):
    try:
        o = subprocess.run(["pdfinfo", P(rel)], capture_output=True, text=True).stdout
        return int(re.search(r"Pages:\s+(\d+)", o).group(1))
    except Exception:
        return len(re.findall(rb"/Type\s*/Page[^s]", open(P(rel), "rb").read()))

README, UG = read("README.md") if os.path.isfile(P("README.md")) else "", read("04_Docs/User_Guide.md") if os.path.isfile(P("04_Docs/User_Guide.md")) else ""
MAN = json.load(open(P("04_Docs/analysis/number_manifest_step8.json"), encoding="utf-8"))
FACTS = json.load(open(P("04_Docs/insights_facts.json"), encoding="utf-8"))
LIVE5 = json.load(open(P("03_PowerBI/validation/step5_live_results.json"), encoding="utf-8-sig"))
LIVE6 = json.load(open(P("03_PowerBI/validation/step6_live_values.json"), encoding="utf-8-sig"))
PBI_MEASURES = LIVE5["measures"]
PBI_TOTAL = dict(zip([c.strip("[]") for c in LIVE5["queries"]["total"]["columns"]], LIVE5["queries"]["total"]["rows"][0]))
DECK = "05_Presentation/Talent_Learning_Deck.pptx"

# ------------------------------------------------------------------------------------------- A. deliverables
S = "A. Deliverables"
need = ["README.md", "LICENSE", ".gitignore", "04_Docs/User_Guide.md", "04_Docs/User_Guide.pdf", "04_Docs/Dashboard_Preview.pdf",
        DECK, "05_Presentation/Talent_Learning_Deck.pdf", "screenshots/star_schema.png", "04_Docs/Insights_and_Recommendations.md"]
missing = [f for f in need if not os.path.isfile(P(f)) or os.path.getsize(P(f)) == 0]
check(S, f"All {len(need)} Step 8 deliverables exist and are non-empty", not missing, str(missing))
shots = sorted(glob.glob(P("screenshots", "0[1-8]_*.png")))
from PIL import Image
sizes = {Image.open(s).size for s in shots}
check(S, "8 page screenshots, all the same size and 16:9 (cropped to the report canvas)",
      len(shots) == 8 and len(sizes) == 1 and abs(list(sizes)[0][0] / list(sizes)[0][1] - 16 / 9) < 0.01, f"{len(shots)} files, sizes={sizes}")
pages_order = json.load(open(P("03_PowerBI/Talent_Learning_Dashboard.Report/definition/pages/pages.json"), encoding="utf-8-sig"))["pageOrder"]
names = [json.load(open(P("03_PowerBI/Talent_Learning_Dashboard.Report/definition/pages", p, "page.json"), encoding="utf-8-sig"))["displayName"] for p in pages_order]
slug = [re.sub(r"[^a-z]+", "_", n.lower()).strip("_") for n in names]
check(S, "Screenshot files are named after the report pages in report order",
      [os.path.basename(s)[3:-4] for s in shots] == slug, str([os.path.basename(s) for s in shots]))
check(S, "User guide PDF is 1-2 pages (as planned in System Design)", 1 <= pdf_pages("04_Docs/User_Guide.pdf") <= 2, str(pdf_pages("04_Docs/User_Guide.pdf")))
check(S, "Dashboard preview PDF has one page per report page", pdf_pages("04_Docs/Dashboard_Preview.pdf") == len(names), str(pdf_pages("04_Docs/Dashboard_Preview.pdf")))
check(S, "Deck PDF has 5 pages", pdf_pages("05_Presentation/Talent_Learning_Deck.pdf") == 5, str(pdf_pages("05_Presentation/Talent_Learning_Deck.pdf")))
lic = read("LICENSE")
check(S, "LICENSE is MIT and states that the dataset is under IBM's ODbL / DbCL terms", lic.startswith("MIT License") and "ODbL" in lic and "DbCL" in lic)

# ------------------------------------------------------------------------------------------- B. links
S = "B. Links"
EXTERNAL_OK = {"https://www.kaggle.com/datasets/pavansubhasht/ibm-hr-analytics-attrition-dataset", "https://github.com/IBM/employee-attrition-aif360",
               "https://www.microsoft.com/power-platform/products/power-bi/desktop", "https://github.com/Kapil-Sandhu"}
def links(md):
    return re.findall(r"!?\[[^\]]*\]\(([^)\s]+)\)", md) + re.findall(r'<img[^>]+src="([^"]+)"', md)
for name, md, base in (("README.md", README, ""), ("User_Guide.md", UG, "04_Docs")):
    bad, ext = [], []
    for l in links(md):
        if l.startswith("http"): (ext.append(l) if l not in EXTERNAL_OK else None); continue
        if not os.path.exists(os.path.normpath(P(base, l.split("#")[0]))): bad.append(l)
    check(S, f"{name}: every relative link / image resolves ({len(links(md))} links)", not bad, str(bad))
    check(S, f"{name}: external links are only the known sources", not ext, str(ext))
def code_spans(md): return re.findall(r"`([^`\n]+)`", re.sub(r"```.*?```", "", md, flags=re.S))
mnames = {m["name"] for m in PBI_MEASURES}
paths = [c for c in code_spans(README + "\n" + UG) if ("/" in c or c.endswith((".pbip", ".md", ".py", ".pdf"))) and not c.startswith(("pip ", "git ")) and "*" not in c]
bad = [c for c in paths if not os.path.exists(P(c.rstrip("/")))]
check(S, f"Every file path in backticks exists ({len(paths)} paths)", not bad, str(bad))
tree = re.search(r"## Repository structure\n\n```\n(.*?)```", README, re.S)
entries = re.findall(r"[├└│ ─]+\s*([^\s][^\s]*?/?)(?:\s{2,}|\s*<-|\n|$)", tree.group(1)) if tree else []
stack, bad, seen = [], [], 0
for line in (tree.group(1).splitlines() if tree else []):
    m = re.match(r"^((?:│   |    )*)(?:├── |└── )(\S+(?: \S+)*?)(?:\s{2,}|\s*<-|$)", line)
    if not m: continue
    depth = len(m.group(1)) // 4; name = m.group(2)
    stack = stack[:depth] + [name.rstrip("/")]; seen += 1
    if not os.path.exists(P(*stack)): bad.append("/".join(stack))
check(S, f"Every entry of the README's repository tree exists ({seen} entries)", seen >= 10 and not bad, str(bad))
unknown_measures = [c for c in code_spans(README) if c in mnames or c.startswith(("Attrition", "Avg", "%"))]
check(S, "Measure names in backticks are real measures", all(c in mnames for c in unknown_measures), str([c for c in unknown_measures if c not in mnames]))

# ------------------------------------------------------------------------------------------- D. DAX reference (before C: verified rows are removed from the number audit)
S = "D. DAX reference"
norm = lambda s: re.sub(r"\s+", " ", (s or "").strip())
rows = re.findall(r"^\| `([^`]+)` \| ([^|]+) \| (.+) \|$", README, re.M)
check(S, f"README measure table lists all {len(PBI_MEASURES)} measures in the model's order", [r[0] for r in rows] == [m["name"] for m in PBI_MEASURES],
      str([r[0] for r in rows][:3]))
desc_bad = [r[0] for r, m in zip(rows, PBI_MEASURES) if norm(r[2]) != norm(m["description"])]
check(S, "Every measure description in the README == the description Power BI reports", not desc_bad, str(desc_bad))
blocks = re.findall(r"\*\*([^*]+)\*\*  \(format `([^`]+)`\)\n```dax\n(.*?)\n```", README, re.S)
bmap = {b[0]: b for b in blocks}
code_bad = [m["name"] for m in PBI_MEASURES if m["name"] not in bmap or norm(bmap[m["name"]][2]) != norm(m["expression"]) or bmap[m["name"]][1] != m["formatString"]]
check(S, f"README DAX code + format for all {len(PBI_MEASURES)} measures == what Power BI reports", len(blocks) == len(PBI_MEASURES) and not code_bad, str(code_bad[:4]))
folders = re.findall(r"<summary><b>DAX - ([^<]+)</b> \((\d+) measures\)</summary>", README)
fcount = {}
for m in PBI_MEASURES: fcount[m["displayFolder"][2:]] = fcount.get(m["displayFolder"][2:], 0) + 1
check(S, "DAX sections == the model's display folders, with the right measure counts", dict((a, int(b)) for a, b in folders) == fcount, str(folders))
README_AUDIT = re.sub(r"```.*?```", " ", README, flags=re.S)                              # code verified above
README_AUDIT = re.sub(r"\(format `[^`]+`\)", " ", README_AUDIT)                          # format strings verified above
README_AUDIT = re.sub(r"<summary>.*?</summary>", " ", README_AUDIT)                        # folder names + counts verified above
README_AUDIT = re.sub(r"<[^>]+>", " ", README_AUDIT)                                      # HTML attributes (width="50%")
README_AUDIT = re.sub(r"^\| `[^`]+` \| [^|]+ \| (.+) \|$", lambda m: "| " + m.group(0).split("|")[2] + " |", README_AUDIT, flags=re.M)  # keep only the value cell

# ------------------------------------------------------------------------------------------- C. number audit
S = "C. Number audit"
NUM = r"\d{1,3}(?:,\d{3})+(?:\.\d+)?|\d+(?:\.\d+)?"
TOK = re.compile(rf"(?<![A-Za-z_\d.,/])(?:(?:{NUM})-(?:{NUM})%?|[+-]?(?:{NUM})(?:%|x| pts|-day)?(?:/(?:{NUM}))?)(?![_\d])")
def r_half(x, dp):
    q = Decimal(1).scaleb(-dp); d = Decimal(repr(float(x)))
    return d.quantize(q, rounding=ROUND_HALF_UP) if d >= 0 else -(-d).quantize(q, rounding=ROUND_HALF_UP)
# justified literals: dimension band labels, rating-scale ends, policy literals of the recommendations
lk = {}
from openpyxl import load_workbook
wb = load_workbook(P("01_Data/lookups/HR_Lookups.xlsx"), data_only=True)
for ws in wb.worksheets:
    for tn in ws.tables:
        rr = [[c.value for c in row] for row in ws[ws.tables[tn].ref]]; lk[tn] = pd.DataFrame(rr[1:], columns=rr[0])
band_tokens = set()
for f in glob.glob(P("03_PowerBI/validation/expected/Dim*Band.csv")):
    for v in pd.read_csv(f).astype(str).values.ravel(): band_tokens |= set(TOK.findall(v))
sat = lk["Lkp_Satisfaction"].iloc[:, 0]; wlb = lk["Lkp_WorkLifeBalance"].iloc[:, 0]
scale = {str(int(sat.min())), str(int(sat.max())), f"{int(sat.min())}-{int(sat.max())}"} | {str(int(wlb.min())), str(int(wlb.max()))}
POLICY = {"90-day": "recommendation 2 (onboarding plan)", "30": "recommendation 4 (follow-up within 30 days)"}
ALLOWED = band_tokens | scale | set(POLICY) | {"1"}       # "1 fact" table, "1 row per employee"
def strip_structure(t):
    t = re.sub(r"(?m)^\s*\d+\.\s", " ", t)                          # ordered-list markers
    t = re.sub(r"(?m)^#+\s+\d+\.\s", " ", t)                        # numbered headings
    t = re.sub(r"(?m)^\| \d \| ", "| ", t)                          # step number in the validation table
    t = re.sub(r"\bSteps? \d+(?:-\d+)?", " ", t)                    # "Step 7", "Steps 1-7"
    t = re.sub(r"validation_report_step\d|step\d", " ", t)
    return t
def audit(doc, text):
    logged = {m["text"] for m in MAN if m["doc"] == doc}
    toks = TOK.findall(strip_structure(text))
    unexplained = sorted({t for t in toks if t not in logged and t.lstrip("+") not in logged and t not in ALLOWED
                          and not any(t in L.split() for L in logged)})
    return toks, unexplained, logged
for doc, text in (("README.md", README_AUDIT), ("04_Docs/User_Guide.md", UG)):
    toks, un, logged = audit(doc, text)
    check(S, f"{doc}: every one of its {len(toks)} numbers is logged with a source or is a justified literal", not un, str(un[:12]))
    miss = [t for t in logged if t not in text]
    check(S, f"{doc}: every logged number appears in it", not miss, str(miss[:6]))
# deck text from the PPTX itself (slides + notes + chart title), skipping the numbered circles
NS = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main", "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
      "c": "http://schemas.openxmlformats.org/drawingml/2006/chart"}
z = zipfile.ZipFile(P(DECK))
slides = sorted([n for n in z.namelist() if re.fullmatch(r"ppt/slides/slide\d+\.xml", n)], key=lambda n: int(re.search(r"(\d+)", n).group(1)))
deck_text, slide_titles, shape_texts = [], [], []
for sname in slides:
    root = ET.fromstring(z.read(sname)); title = ""
    for sp in root.iter(f"{{{NS['p']}}}sp"):
        nv = sp.find(".//p:cNvPr", NS); nm = nv.get("name", "") if nv is not None else ""
        ph = sp.find(".//p:nvPr/p:ph", NS)
        txt = " ".join("".join(t.text or "" for t in para.iter(f"{{{NS['a']}}}t")) for para in sp.iter(f"{{{NS['a']}}}p")).strip()
        if ph is not None and ph.get("type") == "title": title = txt
        if txt and not nm.endswith(" number"): deck_text.append(txt)
        if txt: shape_texts.append((sname, nm, sp))
    slide_titles.append(title)
notes = [n for n in z.namelist() if re.fullmatch(r"ppt/notesSlides/notesSlide\d+\.xml", n)]
notes_text = {n: " ".join(t.text or "" for t in ET.fromstring(z.read(n)).iter(f"{{{NS['a']}}}t")) for n in notes}
charts = [n for n in z.namelist() if re.fullmatch(r"ppt/charts/chart\d+\.xml", n)]
chart_titles = [" ".join(t.text or "" for t in ET.fromstring(z.read(c)).iter(f"{{{NS['a']}}}t")) for c in charts]
DECK_ALL = "\n".join(deck_text + list(notes_text.values()) + chart_titles)
toks, un, logged = audit("05_Presentation/deck_content.json", DECK_ALL)
check(S, f"Deck (slides, speaker notes, chart): every one of its {len(toks)} numbers is logged or a justified literal", not un, str(un[:12]))
check(S, "Deck: every logged number appears in the PPTX", not [t for t in logged if t not in DECK_ALL], str([t for t in logged if t not in DECK_ALL][:6]))
# re-check each manifest entry against an independent source
def fval(path):
    v = FACTS
    for k in path.split("."): v = v[k]
    return v
def pbi_fmt(v, fs):
    if fs == "#,##0": return f"{r_half(v, 0):,}"
    if fs == "0.0%": return f"{r_half(v * 100, 1)}%"
    if fs in ("0.0", "0.00"): return f"{r_half(v, len(fs) - 2)}"
    if fs == '0.0"x"': return f"{r_half(v, 1)}x"
    if fs == "+0.0%;-0.0%;0.0%":
        t = r_half(v * 100, 1); return "0.0%" if t == 0 else (f"+{t}%" if t > 0 else f"{t}%")
fs_of = {m["name"]: m["formatString"] for m in PBI_MEASURES}
rd = P("03_PowerBI/Talent_Learning_Dashboard.Report/definition")
recount = {
    "pages": len([d for d in os.listdir(os.path.join(rd, "pages")) if os.path.isdir(os.path.join(rd, "pages", d))]),
    "visuals": sum(len(fs) for _, _, fs in os.walk(os.path.join(rd, "pages")) if "visual.json" in fs),
    "measures": len(re.findall(r"(?m)^\tmeasure ", read("03_PowerBI/Talent_Learning_Dashboard.SemanticModel/definition/tables/_Measures.tmdl"))),
    "folders": len({m["displayFolder"] for m in PBI_MEASURES}),
    "queries": len(re.findall(r"(?m)^expression ", read("03_PowerBI/Talent_Learning_Dashboard.SemanticModel/definition/expressions.tmdl")))
               + sum(1 for f in glob.glob(P("03_PowerBI/Talent_Learning_Dashboard.SemanticModel/definition/tables/*.tmdl"))
                     if re.search(r"(?m)^\tpartition .* = m$", read(os.path.relpath(f, ROOT))) and not re.search(r"source = #table\(type table \[[^\]]*\], \{\}\)", read(os.path.relpath(f, ROOT)))),
    "csv_rows": len(pd.read_csv(P("01_Data/raw/WA_Fn-UseC_-HR-Employee-Attrition.csv"))),
    "csv_cols": pd.read_csv(P("01_Data/raw/WA_Fn-UseC_-HR-Employee-Attrition.csv")).shape[1],
    "relationships": len(re.findall(r"(?m)^relationship ", read("03_PowerBI/Talent_Learning_Dashboard.SemanticModel/definition/relationships.tmdl"))),
    "dimensions": len([f for f in glob.glob(P("03_PowerBI/Talent_Learning_Dashboard.SemanticModel/definition/tables/Dim*.tmdl"))]),
    "qa_checks": len(next(v for k, v in LIVE6["queries"].items() if k.endswith("|qa"))["rows"]),
}
qa_rows = next(v for k, v in LIVE6["queries"].items() if k.endswith("|qa"))
qa_all_pass = all(dict(zip(qa_rows["columns"], r)).get("QA_Checks[Result]") == "PASS" for r in qa_rows["rows"])
placeholders = [os.path.basename(f) for f in glob.glob(P("03_PowerBI/Talent_Learning_Dashboard.SemanticModel/definition/tables/*.tmdl"))
                if re.search(r"source = #table\(type table \[[^\]]*\], \{\}\)", read(os.path.relpath(f, ROOT)))]
check(S, "Only one M partition is not ETL: the empty _Measures placeholder table that holds the measures", placeholders == ["_Measures.tmdl"], str(placeholders))
check(S, "Power Query ETL count agrees two ways (.pq files == TMDL shared expressions + M table partitions)",
      recount["queries"] == len(glob.glob(P("03_PowerBI/PowerQuery/*.pq"))), f"{recount['queries']} vs {len(glob.glob(P('03_PowerBI/PowerQuery/*.pq')))}")
def report_tally(step):
    f = {1: "01_Data/validation_report_step1.txt", 2: "02_Excel_Profiling/validation_report_step2.txt", 7: "04_Docs/analysis/validation_report_step7.txt",
         8: "04_Docs/analysis/validation_report_step8.txt"}.get(
        step, f"03_PowerBI/validation/validation_report_step{step}.txt")
    t = read(f); body = t.split("NEGATIVE CONTROLS")[0]
    return len(re.findall(r"(?m)^\[PASS\]", body)), len(re.findall(r"(?m)^\[FAIL\]", body)), t
bad = []
for m in MAN:
    kind, _, key = m["source"].partition(":")
    if kind == "facts":
        if key.startswith("overtime_by_role["):
            v = sum(x["rate_ot"] > x["rate_no_ot"] for x in FACTS["overtime_by_role"].values()) if "rate_ot >" in key else len(FACTS["overtime_by_role"])
            ok = m["value"] == v and m["text"] == str(v)
        else:
            v = fval(key); t = m["text"]
            if m["value"] != v: ok = False
            elif t.endswith(" pts"): ok = t == ("+" if v >= 0 else "") + f"{r_half(v, 1)} pts"
            elif t.endswith("x"): ok = t == f"{r_half(v, 1)}x"
            elif t.endswith("%") and abs(v) > 1: ok = t == pbi_fmt(v / 100, "+0.0%;-0.0%;0.0%")     # points shown in the measure's format
            elif t.endswith("%"): ok = t == f"{r_half(v * 100, 1)}%"
            elif isinstance(v, int) or float(v).is_integer(): ok = t == f"{int(v):,}"
            else: ok = t == f"{r_half(v, 1)}"
    elif kind == "measure":
        pv = PBI_TOTAL[key]; ok = abs(pv - m["value"]) < 1e-9 and m["text"] == pbi_fmt(pv, fs_of[key])   # Power BI's own value
    elif kind == "project":
        ok = recount[key] == m["value"] and m["text"] == f"{recount[key]:,}"
    elif kind == "validation":
        if key == "total":
            tot = sum(report_tally(s)[0] for s in range(1, 8)); ok = m["value"] == tot and m["text"] == f"{tot:,}" and all(report_tally(s)[1] == 0 for s in range(1, 8))
        elif key == "7-negative":
            mm = re.search(r"NEGATIVE CONTROLS: (\d+)/(\d+)", report_tally(7)[2]); n_c = len(re.findall(r"(?m)^  CAUGHT ", report_tally(7)[2]))
            ok = mm and mm.group(1) == mm.group(2) and int(mm.group(1)) == n_c == m["value"]
        else:
            p_, f_, _ = report_tally(int(key)); ok = f_ == 0 and m["text"] == f"{p_}/{p_}"
    else:
        ok = False
    if not ok: bad.append(f"{m['source']}: '{m['text']}'")
check(S, f"All {len(MAN)} logged numbers re-checked against an independent source (facts / Power BI / file recount / report tally)", not bad, "; ".join(bad[:5]))
check(S, "Data Quality page (Power BI): all QA checks PASS, so the README's 'every row should say PASS' is true", qa_all_pass)

# regenerate the documents from the current sources: the committed files must be exactly what the generator writes
tmpd = tempfile.mkdtemp()
try:
    shutil.copytree(ROOT, os.path.join(tmpd, "r"), ignore=lambda dp, n: [x for x in n if x in ("__pycache__", ".git") or (x == "raw" and dp.endswith("screenshots"))])
    gen = subprocess.run([sys.executable, "-W", "ignore", os.path.join(tmpd, "r", "04_Docs", "build", "build_docs.py"), os.path.join(tmpd, "r")],
                         capture_output=True, text=True)
    same = {f: gen.returncode == 0 and open(os.path.join(tmpd, "r", f), "rb").read() == open(P(f), "rb").read()
            for f in ("README.md", "04_Docs/User_Guide.md", "05_Presentation/deck_content.json", "04_Docs/analysis/number_manifest_step8.json")}
finally:
    shutil.rmtree(tmpd, ignore_errors=True)
check(S, "README, user guide, deck content and manifest are exactly what build_docs.py generates from the current sources (no manual edits)",
      all(same.values()), str({k: v for k, v in same.items() if not v}) + (gen.stderr[-300:] if gen.returncode else ""))
def strings(o):
    if isinstance(o, str): yield o
    elif isinstance(o, dict): yield from (x for v in o.values() for x in strings(v))
    elif isinstance(o, list): yield from (x for v in o for x in strings(v))
DCJ = json.load(open(P("05_Presentation/deck_content.json"), encoding="utf-8"))
skip = {DCJ["model"]["image"]} | {d["label"] for d in DCJ["findings"]["chart"]} | {d["path"] for d in DCJ["findings"]["chart"]}
norm_ws = lambda t: re.sub(r"\s+", " ", t).strip()
pptx_all = norm_ws(" ".join(deck_text) + " " + " ".join(notes_text.values()) + " " + " ".join(chart_titles))
absent = [t for t in strings(DCJ) if t not in skip and norm_ws(t) not in pptx_all]
check(S, "Every text in deck_content.json appears verbatim in the PPTX (the deck was rebuilt after the last content change)", not absent, str(absent[:4]))

# ------------------------------------------------------------------------------------------- E. deck
S = "E. Deck"
check(S, "5 slides, each with a title", len(slides) == 5 and all(slide_titles), str(slide_titles))
rels_ok = all(any(f"notesSlide" in t for t in [z.read(f"ppt/slides/_rels/{os.path.basename(s)}.rels").decode()]) for s in slides)
check(S, "Every slide has speaker notes", rels_ok and len(notes) == 5 and all(len(t.strip()) > 40 for t in notes_text.values()), str({k: len(v) for k, v in notes_text.items()}))
DC = json.load(open(P("05_Presentation/deck_content.json"), encoding="utf-8"))
cx = ET.fromstring(z.read(charts[0])) if charts else None
sers = {}
if cx is not None:
    for ser in cx.iter(f"{{{NS['c']}}}ser"):
        nm = "".join(t.text for t in ser.find("c:tx", NS).iter(f"{{{NS['c']}}}v"))
        sers[nm] = [float(v.text) for v in ser.find("c:val", NS).iter(f"{{{NS['c']}}}v")]
        cats = [v.text for v in ser.find("c:cat", NS).iter(f"{{{NS['c']}}}v")]
exp_vals = [fval(d["path"]) for d in DC["findings"]["chart"]]
got = [a + b for a, b in zip(sers.get("Group", []), sers.get("Company", []))]
check(S, "Chart values in the PPTX == insights_facts.json (each bar), company bar separate",
      len(charts) == 1 and len(got) == len(exp_vals) and all(abs(a - b) < 1e-12 for a, b in zip(got, exp_vals))
      and sum(1 for v in sers.get("Company", []) if v) == 1, f"{got} vs {exp_vals}")
check(S, "Chart labels show 1 decimal percent and hide the zero half of each stacked bar", b'formatCode="0.0%;;;"' in z.read(charts[0]) if charts else False)
check(S, "Star-schema image is embedded", any(n.startswith("ppt/media/") and n.endswith(".png") for n in z.namelist()))
EMU = 914400
off = []
for sname, nm, sp in shape_texts:
    x = sp.find(".//a:xfrm", NS)
    if x is None: continue
    o, e = x.find("a:off", NS), x.find("a:ext", NS)
    if int(o.get("x")) + int(e.get("cx")) > 13.333 * EMU + 10 or int(o.get("y")) + int(e.get("cy")) > 7.5 * EMU + 10: off.append(f"{sname}:{nm}")
check(S, "No text box extends past the slide edge", not off, str(off))
check(S, "Deck title promises match the plan (problem -> data -> model -> findings -> recommendations)",
      [s.split(" ")[0] for s in slide_titles] == [DC["title"]["headline"].split(" ")[0], "From", "A", "Overtime", "Four"], str(slide_titles))

# ------------------------------------------------------------------------------------------- F. consistency across all Markdown docs
S = "F. Consistency"
STEP_OF_FILE = {"03_PowerBI/PowerQuery/README_PowerQuery_Setup.md": 3, "03_PowerBI/model/Model_Design.md": 4, "03_PowerBI/model/DAX_Measures.md": 5,
                "03_PowerBI/report/Report_Guide.md": 6}
mds = sorted(os.path.relpath(f, ROOT).replace("\\", "/") for f in glob.glob(P("**", "*.md"), recursive=True)
             if "incoming" not in f and "backup_" not in f)
bad_claims, n_claims = [], 0
count_words = {"pages": "pages", "visuals": "visuals", "measures": "measures", "queries": "queries", "QA checks": "qa_checks", "relationships": "relationships",
               "display folders": "folders"}
slide_n = len(slides)
for rel in mds:
    for ln, line in enumerate(read(rel).splitlines(), 1):
        for m in re.finditer(r"(?<![\d/])(\d+)/(\d+)(?=\s*(?:checks|passed|\+|incl|:|\)|\]| planted| deliberately))", line):
            sm = re.search(r"step(\d)|- \[[x ]\] (\d)\.|^\| (\d) \||Step (\d) validator", line)
            step = int(next(g for g in sm.groups() if g)) if sm else STEP_OF_FILE.get(rel)
            if step is None: continue
            if re.match(r"\s*(?:deliberately\s+)?(?:planted|caught)", line[m.end():]):   # a negative-control count ("25/25 planted", "18/18 caught")
                n_claims += 1
                nm_ = re.search(r"NEGATIVE CONTROLS: (\d+)/(\d+)", report_tally(step)[2])
                if not (nm_ and nm_.group(1) == nm_.group(2) == m.group(1) == m.group(2)):
                    bad_claims.append(f"{rel}:{ln} says {m.group(0)} planted faults caught for step {step}, report says {nm_.group(0) if nm_ else 'none'}")
                continue
            n_claims += 1
            p_, f_, _ = report_tally(step)
            if not (int(m.group(1)) == int(m.group(2)) == p_ and f_ == 0): bad_claims.append(f"{rel}:{ln} says {m.group(0)} for step {step}, report has {p_}")
        for m in re.finditer(r"(?<![\d–\-/])\b(\d+) (pages|visuals|measures|queries|QA checks|relationships|display folders)\b", "" if "<summary>" in line else line):
            if rel == "System Design.md" and ln < 190 and m.group(2) == "pages" and "guide" in line.lower(): continue   # "user guide (1-2 pages)"
            n_claims += 1
            if int(m.group(1)) != recount[count_words[m.group(2)]]: bad_claims.append(f"{rel}:{ln} says {m.group(0)}, project has {recount[count_words[m.group(2)]]}")
        for m in re.finditer(r"\b(\d+)-slide\b", line):
            n_claims += 1
            if int(m.group(1)) != slide_n: bad_claims.append(f"{rel}:{ln} says {m.group(0)}, deck has {slide_n}")
check(S, f"Every check tally and project count claimed in the {len(mds)} Markdown docs is current ({n_claims} claims)", not bad_claims, "; ".join(bad_claims[:5]))
sd = read("System Design.md")
check(S, "System Design progress checklist: Steps 1-7 ticked and each links an existing report",
      all(re.search(rf"- \[x\] {s}\.", sd) for s in range(1, 8)) and all(os.path.isfile(P(p)) for p in re.findall(r"`([^`]+validation_report_step\d\.txt)`", sd)))

# ------------------------------------------------------------------------------------------- G. git readiness
S = "G. Git readiness"
tmp = tempfile.mkdtemp()
try:
    env = dict(os.environ, GIT_DIR=os.path.join(tmp, "g.git"), GIT_WORK_TREE=os.path.abspath(ROOT))
    subprocess.run(["git", "init", "-q"], env=env, check=True, capture_output=True)
    subprocess.run(["git", "config", "core.quotepath", "off"], env=env, check=True)
    subprocess.run(["git", "config", "core.autocrlf", "false"], env=env, check=True)
    ls = subprocess.run(["git", "add", "--dry-run", "-A"], env=env, capture_output=True, text=True, check=True).stdout
    files = [re.sub(r"^add '(.*)'$", r"\1", l) for l in ls.splitlines() if l.startswith("add ")]
finally:
    shutil.rmtree(tmp, ignore_errors=True)
forbidden = [f for f in files if re.search(r"(^|/)\.pbi/(cache\.abf|localSettings\.json)$|/incoming/|/backup_|__pycache__|\.pyc$|screenshots/raw/|/schemas/|~\$|_render_tmp", f)]
check(S, f"git would commit {len(files)} files and none of them is a Power BI cache, local setting or leftover", files and not forbidden, str(forbidden[:6]))
big = [f for f in files if os.path.getsize(P(f)) > 50 * 1024 * 1024]
total = sum(os.path.getsize(P(f)) for f in files)
check(S, f"No file over 50 MB (GitHub warns at 50 MB); repo size {total / 1e6:.1f} MB", not big, str(big))
must = ["README.md", "LICENSE", "03_PowerBI/Talent_Learning_Dashboard.pbip", "03_PowerBI/Talent_Learning_Dashboard.SemanticModel/definition/model.tmdl",
        "03_PowerBI/Talent_Learning_Dashboard.Report/definition/report.json", "01_Data/raw/WA_Fn-UseC_-HR-Employee-Attrition.csv", DECK,
        "screenshots/01_overview.png", "04_Docs/User_Guide.pdf"]
check(S, "The files needed to open and review the project are all included", all(m in files for m in must), str([m for m in must if m not in files]))
SECRET = re.compile(r"gh[pousr]_[A-Za-z0-9]{20,}|AKIA[0-9A-Z]{16}|-----BEGIN [A-Z ]*PRIVATE KEY|xox[baprs]-|password\s*[:=]\s*\S+|api[_-]?key\s*[:=]\s*\S+", re.I)
PERSONAL = re.compile(r"[A-Za-z0-9._%+-]+@(?:gmail|uni-bonn|outlook|hotmail|yahoo)\.[a-z.]+|C:\\\\?Users\\\\?(?!you\\)[A-Za-z]+", re.I)
hits = []
for f in files:
    if os.path.splitext(f)[1].lower() in (".png", ".pdf", ".pptx", ".xlsx", ".abf"): continue
    try: t = open(P(f), encoding="utf-8", errors="ignore").read()
    except Exception: continue
    for rx, kind in ((SECRET, "secret"), (PERSONAL, "personal")):
        for mm in rx.finditer(t): hits.append(f"{kind}:{f}:{mm.group(0)[:3]}***")   # masked - never repeat a secret in the report
check(S, "No secrets (tokens, keys, passwords) or personal e-mail / user-folder paths in any committed text file", not hits, str(hits[:5]))
pq = read("03_PowerBI/Talent_Learning_Dashboard.SemanticModel/definition/expressions.tmdl")
check(S, "README tells a new user how to point ProjectFolder at their own copy (the parameter exists in the model)",
      "expression ProjectFolder" in pq and "Edit parameters" in README and "ProjectFolder" in README)
gi = read(".gitignore")
check(S, ".gitignore excludes the Power BI cache and local settings everywhere", "**/.pbi/cache.abf" in gi and "**/.pbi/localSettings.json" in gi)

passed, total_ = sum(res), len(res)
text = f"STEP 8 VALIDATION: {passed}/{total_} checks passed\n\n" + "\n".join(out)
print(text)
open(REPORT, "w", encoding="utf-8").write(text + "\n")
sys.exit(0 if passed == total_ else 1)
