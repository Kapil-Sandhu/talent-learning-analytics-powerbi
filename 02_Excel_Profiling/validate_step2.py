"""
Step 2 validation - independent re-computation of every number in the Excel workbooks.
Recomputes from the raw CSV with pandas and compares with Excel's calculated values.
Requires: pandas, openpyxl. Workbooks must have been opened/recalculated (Excel saves cached values).
Run:  python validate_step2.py <csv> <HR_Lookups.xlsx> <HR_Data_Profiling.xlsx> <report.txt>
"""
import sys, re
import pandas as pd
from openpyxl import load_workbook

a = sys.argv[1:5] or ["01_Data/raw/WA_Fn-UseC_-HR-Employee-Attrition.csv", "01_Data/lookups/HR_Lookups.xlsx",
                     "02_Excel_Profiling/HR_Data_Profiling.xlsx", "02_Excel_Profiling/validation_report_step2.txt"]
CSV, LKP, PRF, REPORT = a
df = pd.read_csv(CSV, encoding="utf-8-sig")
N = len(df)
res, out = [], []

def check(section, name, ok, detail=""):
    res.append(ok)
    out.append(f"[{'PASS' if ok else 'FAIL'}] {section}: {name}" + (f" - {detail}" if detail else ""))

def close(a, b, tol=1e-9):
    try:
        return abs(float(a) - float(b)) <= tol
    except (TypeError, ValueError):
        return False

lk_v = load_workbook(LKP, data_only=True)
pf_v = load_workbook(PRF, data_only=True)
pf_f = load_workbook(PRF)  # formulas

# ------------------------------------------------------------------ A. Lookup workbook
S = "A. Lookups"
expected_tables = {"Lkp_Education", "Lkp_Satisfaction", "Lkp_WorkLifeBalance", "Lkp_PerformanceRating",
                   "Lkp_BusinessTravel", "Lkp_AgeBand", "Lkp_TenureBand", "Lkp_IncomeBand", "Lkp_TrainingBand"}
found = {}
for ws in lk_v.worksheets:
    for tname, tbl in ws.tables.items():
        rows = list(ws[ws.tables[tname].ref])
        hdr = [c.value for c in rows[0]]
        found[tname] = pd.DataFrame([[c.value for c in r] for r in rows[1:]], columns=hdr)
check(S, "All 9 lookup tables exist with correct names", set(found) == expected_tables,
      f"missing={expected_tables - set(found)} extra={set(found) - expected_tables}")

def codes_cover(tname, codecol, datacols):
    codes = set(found[tname][codecol])
    for dc in datacols:
        vals = set(df[dc].unique())
        check(S, f"{tname} covers every {dc} value", vals <= codes, f"unmapped={vals - codes}")
    check(S, f"{tname} codes unique", found[tname][codecol].is_unique)
    check(S, f"{tname} labels non-empty", found[tname].iloc[:, 1].notna().all() and (found[tname].iloc[:, 1].astype(str).str.strip() != "").all())

codes_cover("Lkp_Education", "EducationCode", ["Education"])
codes_cover("Lkp_Satisfaction", "RatingCode", ["JobSatisfaction", "EnvironmentSatisfaction", "RelationshipSatisfaction", "JobInvolvement"])
codes_cover("Lkp_WorkLifeBalance", "WLBCode", ["WorkLifeBalance"])
codes_cover("Lkp_PerformanceRating", "PerformanceCode", ["PerformanceRating"])
codes_cover("Lkp_BusinessTravel", "BusinessTravelCode", ["BusinessTravel"])

# documented labels (Kaggle dataset description)
doc = {"Lkp_Education": {1: "Below College", 2: "College", 3: "Bachelor", 4: "Master", 5: "Doctor"},
       "Lkp_Satisfaction": {1: "Low", 2: "Medium", 3: "High", 4: "Very High"},
       "Lkp_WorkLifeBalance": {1: "Bad", 2: "Good", 3: "Better", 4: "Best"},
       "Lkp_PerformanceRating": {1: "Low", 2: "Good", 3: "Excellent", 4: "Outstanding"}}
for tname, m in doc.items():
    got = dict(zip(found[tname].iloc[:, 0], found[tname].iloc[:, 1]))
    check(S, f"{tname} labels match dataset documentation", got == m, f"got={got}")

bands = {"Lkp_AgeBand": "Age", "Lkp_TenureBand": "YearsAtCompany", "Lkp_IncomeBand": "MonthlyIncome",
         "Lkp_TrainingBand": "TrainingTimesLastYear"}
for tname, col in bands.items():
    b = found[tname]
    lo, hi = b.iloc[:, 1].astype(int).tolist(), b.iloc[:, 2].astype(int).tolist()
    contiguous = all(lo[i + 1] == hi[i] + 1 for i in range(len(lo) - 1)) and all(l <= h for l, h in zip(lo, hi))
    check(S, f"{tname} bands contiguous, no gaps/overlaps", contiguous, f"bounds={list(zip(lo, hi))}")
    check(S, f"{tname} covers {col} range {df[col].min()}-{df[col].max()}",
          lo[0] <= df[col].min() and hi[-1] >= df[col].max())
    per_emp = df[col].apply(lambda v: sum(l <= v <= h for l, h in zip(lo, hi)))
    check(S, f"{tname}: every employee in exactly one band", (per_emp == 1).all(),
          f"{(per_emp != 1).sum()} employees not in exactly one band")
    check(S, f"{tname} SortOrder is 1..n", b["SortOrder"].tolist() == list(range(1, len(b) + 1)))

# ------------------------------------------------------------------ B. Raw_Data integrity
S = "B. Raw_Data"
ws = pf_v["Raw_Data"]
check(S, "tbl_HR table exists", "tbl_HR" in ws.tables)
check(S, "tbl_HR range = A1:AI1471", ws.tables["tbl_HR"].ref == f"A1:AI{N + 1}", ws.tables["tbl_HR"].ref)
hdr = [c.value for c in ws[1]]
check(S, "Headers identical to CSV (incl. 'Age' without BOM)", hdr == list(df.columns), str(hdr[:3]))
mismatch = 0
for i, row in enumerate(ws.iter_rows(min_row=2, max_row=N + 1, values_only=True)):
    for j, v in enumerate(row):
        if str(v) != str(df.iat[i, j]):
            mismatch += 1
check(S, f"All {N * len(df.columns):,} cells identical to CSV", mismatch == 0, f"{mismatch} mismatches")

# ------------------------------------------------------------------ C. KPI_Summary
S = "C. KPI_Summary"
k = pf_v["KPI_Summary"]
att = df.Attrition.eq("Yes")
ot = df.OverTime.eq("Yes")
exp = {
    "Headcount": N, "Leavers": att.sum(), "Active Employees": N - att.sum(), "Attrition Rate": att.mean(),
    "Avg Tenure (years)": df.YearsAtCompany.mean(), "Median Tenure (years)": df.YearsAtCompany.median(),
    "Avg Tenure - Leavers (years)": df.loc[att, "YearsAtCompany"].mean(), "Avg Age": df.Age.mean(),
    "Avg Monthly Income": df.MonthlyIncome.mean(), "Total Trainings Last Year": df.TrainingTimesLastYear.sum(),
    "Training per Employee": df.TrainingTimesLastYear.mean(),
    "% Zero Training": df.TrainingTimesLastYear.eq(0).mean(), "Overtime %": ot.mean(),
    "Attrition Rate - Overtime": att[ot].mean(), "Attrition Rate - No Overtime": att[~ot].mean(),
    "Avg Job Satisfaction (1-4)": df.JobSatisfaction.mean(), "% Low Job Satisfaction": df.JobSatisfaction.eq(1).mean(),
    "Avg Years Since Last Promotion": df.YearsSinceLastPromotion.mean(), "Avg Performance Rating": df.PerformanceRating.mean(),
}
seen = {}
for r in range(5, k.max_row + 1):
    if k.cell(row=r, column=1).value:
        seen[k.cell(row=r, column=1).value] = k.cell(row=r, column=2).value
check(S, f"All {len(exp)} KPIs present", set(exp) <= set(seen), f"missing={set(exp) - set(seen)}")
for name, e in exp.items():
    check(S, f"{name} = {e:.6g}", close(seen.get(name), e), f"excel={seen.get(name)}")
kf = pf_f["KPI_Summary"]
hard = [kf.cell(row=r, column=2).coordinate for r in range(5, kf.max_row + 1)
        if kf.cell(row=r, column=2).value is not None and not str(kf.cell(row=r, column=2).value).startswith("=")]
check(S, "Every KPI value is a formula (no hardcoded results)", not hard, str(hard))

# ------------------------------------------------------------------ D. Profile blocks
S = "D. Profiles"
band_defs = {k_: list(zip(found[k_].iloc[:, 0], found[k_].iloc[:, 1], found[k_].iloc[:, 2])) for k_ in bands}
field_band = {"Age": "Lkp_AgeBand", "YearsAtCompany": "Lkp_TenureBand", "MonthlyIncome": "Lkp_IncomeBand",
              "TrainingTimesLastYear": "Lkp_TrainingBand"}

def parse_blocks(sheet):
    """Find blocks: header row with 'Category' in A; read rows until 'Total'. Field taken from formula."""
    wv, wf = pf_v[sheet], pf_f[sheet]
    blocks, r = [], 1
    while r <= wv.max_row:
        if wv.cell(row=r, column=1).value == "Category" and wv.cell(row=r, column=4).value == "Headcount":
            title = wv.cell(row=r - 1, column=1).value
            if title and str(title).startswith("Bands mirror") or (title and wv.cell(row=r - 1, column=1).font.italic):
                title = wv.cell(row=r - 2, column=1).value
            rows, rr = [], r + 1
            while wv.cell(row=rr, column=1).value != "Total":
                rows.append(rr); rr += 1
            f = wf.cell(row=rows[0], column=4).value
            field = re.search(r"tbl_HR\[(\w+)\]", f).group(1)
            kind = "band" if '">="' in f else ("code" if "$B" in f else "text")
            blocks.append((title, field, kind, rows, rr))
            r = rr
        r += 1
    return blocks

total_blocks = 0
for sheet in ["Attrition_Profile", "Training_Profile", "Satisfaction_Profile"]:
    wv, wf = pf_v[sheet], pf_f[sheet]
    for title, field, kind, rows, trow in parse_blocks(sheet):
        total_blocks += 1
        bad = []
        labels_lohi = []
        for r in rows:
            a_, b_, c_ = (wv.cell(row=r, column=x).value for x in (1, 2, 3))
            if kind == "text":
                m = df[field] == a_
            elif kind == "code":
                m = df[field] == b_
            else:
                m = (df[field] >= b_) & (df[field] <= c_)
                labels_lohi.append((a_, b_, c_))
            hc, lv = int(m.sum()), int((m & att).sum())
            rate = (lv / hc) if hc else "n/a"
            got = [wv.cell(row=r, column=x).value for x in (4, 5, 6, 7, 8)]
            exp_row = [hc, lv, rate, (rate - att.mean()) if hc else "n/a", hc / N]
            for g, e in zip(got, exp_row):
                if not (g == e if isinstance(e, str) else close(g, e)):
                    bad.append(f"r{r}:{a_} got={got} exp={exp_row}")
                    break
            for x in range(4, 9):
                if not str(wf.cell(row=r, column=x).value).startswith("="):
                    bad.append(f"{wf.cell(row=r, column=x).coordinate} hardcoded")
        check(S, f"{sheet} / {title} ({len(rows)} rows) match pandas", not bad, "; ".join(bad[:3]))
        tot = wv.cell(row=trow, column=4).value
        if kind == "text" and field in ("Department", "JobRole", "OverTime", "Gender", "MaritalStatus", "BusinessTravel") \
                or kind == "band" or field in ("JobLevel", "StockOptionLevel", "TrainingTimesLastYear") or kind == "code":
            check(S, f"{sheet} / {title} total = {N}", tot == N, f"total={tot}")
        if kind == "band":
            check(S, f"{sheet} / {title} band bounds identical to {field_band[field]}",
                  labels_lohi == band_defs[field_band[field]], f"{labels_lohi} vs {band_defs[field_band[field]]}")
        if kind == "code" and field in ("JobSatisfaction", "EnvironmentSatisfaction", "RelationshipSatisfaction",
                                        "JobInvolvement", "WorkLifeBalance", "PerformanceRating", "Education"):
            lk = {"WorkLifeBalance": "Lkp_WorkLifeBalance", "PerformanceRating": "Lkp_PerformanceRating",
                  "Education": "Lkp_Education"}.get(field, "Lkp_Satisfaction")
            want = [f"{c} - {l}" for c, l in zip(found[lk].iloc[:, 0], found[lk].iloc[:, 1])]
            got_l = [wv.cell(row=r, column=1).value for r in rows]
            check(S, f"{sheet} / {title} labels match {lk}", got_l == want, f"{got_l}")
check(S, "Expected number of profile blocks found (11 + 2 + 7 = 20)", total_blocks == 20, f"found {total_blocks}")

# Training by Department / Job Role tables
wv, wf = pf_v["Training_Profile"], pf_f["Training_Profile"]
for r in range(1, wv.max_row + 1):
    if wv.cell(row=r, column=2).value == "Headcount" and wv.cell(row=r, column=3).value == "Total Trainings":
        field = re.search(r"tbl_HR\[(\w+)\]", wf.cell(row=r + 1, column=2).value).group(1)
        bad, rr = [], r + 1
        while wv.cell(row=rr, column=1).value != "Total":
            v = wv.cell(row=rr, column=1).value
            g = df[df[field] == v]
            exp_row = [len(g), g.TrainingTimesLastYear.sum(), g.TrainingTimesLastYear.mean(),
                       g.TrainingTimesLastYear.eq(0).mean(), g.TrainingTimesLastYear.ge(5).mean(), g.Attrition.eq("Yes").mean()]
            got = [wv.cell(row=rr, column=x).value for x in range(2, 8)]
            if not all(close(a_, b_) for a_, b_ in zip(got, exp_row)):
                bad.append(f"{v}: got={got} exp={exp_row}")
            rr += 1
        got_t = [wv.cell(row=rr, column=x).value for x in range(2, 8)]
        exp_t = [N, df.TrainingTimesLastYear.sum(), df.TrainingTimesLastYear.mean(), df.TrainingTimesLastYear.eq(0).mean(),
                 df.TrainingTimesLastYear.ge(5).mean(), att.mean()]
        if not all(close(a_, b_) for a_, b_ in zip(got_t, exp_t)):
            bad.append(f"Total got={got_t} exp={exp_t}")
        check(S, f"Training_Profile / Training by {field} matches pandas (incl. total)", not bad, "; ".join(bad[:3]))

# ------------------------------------------------------------------ E. Data_Dictionary
S = "E. Data_Dictionary"
wv = pf_v["Data_Dictionary"]
bad, names = [], []
for r in range(5, 5 + len(df.columns)):
    c = wv.cell(row=r, column=2).value
    names.append(c)
    s = df[c]
    numeric = pd.api.types.is_numeric_dtype(s)
    exp_row = [s.min() if numeric else "-", s.max() if numeric else "-", s.nunique(), int(s.isna().sum())]
    got = [wv.cell(row=r, column=x).value for x in (5, 6, 7, 8)]
    if not all((g == e) if isinstance(e, str) else close(g, e) for g, e in zip(got, exp_row)):
        bad.append(f"{c}: got={got} exp={exp_row}")
    if not wv.cell(row=r, column=4).value or not wv.cell(row=r, column=9).value:
        bad.append(f"{c}: missing description/action")
check(S, "All 35 columns listed in CSV order", names == list(df.columns))
check(S, "Min/Max/Distinct/Blanks match pandas; description + action filled", not bad, "; ".join(bad[:3]))
drops = sorted(wv.cell(row=r, column=2).value for r in range(5, 40) if str(wv.cell(row=r, column=9).value).startswith("Drop"))
check(S, "Drop list = 3 constant + 3 undocumented rate columns",
      drops == sorted(["EmployeeCount", "Over18", "StandardHours", "DailyRate", "HourlyRate", "MonthlyRate"]), str(drops))

# ------------------------------------------------------------------ F. Data_Quality
S = "F. Data_Quality"
wv = pf_v["Data_Quality"]
rows = [r for r in range(5, wv.max_row + 1) if wv.cell(row=r, column=5).value in ("PASS", "FAIL")]
fails = [wv.cell(row=r, column=2).value for r in rows if wv.cell(row=r, column=5).value != "PASS"]
check(S, f"All {len(rows)} in-workbook checks PASS", not fails and len(rows) == 19, f"fails={fails}")
summary = [wv.cell(row=r, column=5).value for r in range(1, wv.max_row + 1) if wv.cell(row=r, column=2).value == "Checks passed"]
check(S, "Summary row says ALL PASS", summary == ["ALL PASS"], str(summary))

# ------------------------------------------------------------------ G. Key_Findings
S = "G. Key_Findings"
wv = pf_v["Key_Findings"]
texts = [wv.cell(row=r, column=3).value for r in range(5, 13)]
check(S, "8 findings rendered as text (no errors/blanks)", all(isinstance(t, str) and t and "#" not in t for t in texts))
pct = lambda x: f"{x * 100:.1f}%"
dept_rate = df.groupby("Department").Attrition.apply(lambda s: s.eq("Yes").mean())
role_rate = df.groupby("JobRole").Attrition.apply(lambda s: s.eq("Yes").mean())
expect_snips = [
    [pct(att.mean()), "237", "1,470"],
    [pct(att[ot].mean()), pct(att[~ot].mean())],
    [pct(att[df.YearsAtCompany <= 1].mean()), pct(att[df.YearsAtCompany >= 11].mean())],
    [pct(att[df.JobSatisfaction == 1].mean()), pct(att[df.JobSatisfaction == 4].mean())],
    [dept_rate.idxmax(), pct(dept_rate.max()), role_rate.idxmax(), pct(role_rate.max())],
    [pct(att[df.TrainingTimesLastYear == 0].mean()), pct(att[df.TrainingTimesLastYear >= 5].mean())],
    [pct(att[df.JobLevel == 1].mean()), pct(att[df.StockOptionLevel == 0].mean())],
    [pct(att[df.PerformanceRating == 3].mean()), pct(att[df.PerformanceRating == 4].mean())],
]
for i, (t, snips) in enumerate(zip(texts, expect_snips), start=1):
    miss = [s_ for s_ in snips if s_ not in (t or "")]
    check(S, f"Finding {i} numbers match pandas", not miss, f"missing {miss} in: {t}")

# ------------------------------------------------------------------ H. Formatting / hygiene
S = "H. Formatting"
for name, wb in [("HR_Lookups", load_workbook(LKP)), ("HR_Data_Profiling", pf_f)]:
    non_arial = 0
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if c.value is not None and c.font is not None and c.font.name != "Arial":
                    non_arial += 1
    check(S, f"{name}: every populated cell uses Arial", non_arial == 0, f"{non_arial} cells")
errs = 0
for ws in pf_v.worksheets:
    for row in ws.iter_rows(values_only=True):
        errs += sum(1 for v in row if isinstance(v, str) and v.startswith("#") and v.endswith(("!", "?", "A")))
check(S, "No error values (#DIV/0!, #N/A, #NAME?, #REF!, #VALUE!) in profiling workbook", errs == 0, f"{errs}")
check(S, "Sheet order as designed", pf_f.sheetnames == ["README", "KPI_Summary", "Attrition_Profile", "Training_Profile",
      "Satisfaction_Profile", "Data_Dictionary", "Data_Quality", "Key_Findings", "Raw_Data"], str(pf_f.sheetnames))
import zipfile
charts = len([n_ for n_ in zipfile.ZipFile(PRF).namelist() if re.match(r"xl/charts/chart\d+\.xml$", n_)])
check(S, "Charts present (7)", charts == 7, f"{charts}")

passed, total = sum(res), len(res)
head = f"STEP 2 VALIDATION: {passed}/{total} checks passed"
text = head + "\n\n" + "\n".join(out)
print(text)
open(REPORT, "w", encoding="utf-8").write(text + "\n")
sys.exit(0 if passed == total else 1)
