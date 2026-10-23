"""
Step 2 - Build the Excel profiling workbook and the lookup workbook.
Talent & Learning Analytics Dashboard (IBM HR Attrition dataset, 1,470 employees)

Outputs:
  01_Data/lookups/HR_Lookups.xlsx            - code->label and band lookup tables (Power Query input)
  02_Excel_Profiling/HR_Data_Profiling.xlsx  - formula-driven profiling of the raw data

Every number in the profiling workbook is an Excel formula over table tbl_HR (sheet Raw_Data),
so it recalculates if the raw data changes.
"""
import sys
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.formatting.rule import ColorScaleRule, CellIsRule
from openpyxl.chart import BarChart, Reference
from openpyxl.utils import get_column_letter

# Run from the project root:  python 02_Excel_Profiling/build_step2_workbooks.py
CSV = sys.argv[1] if len(sys.argv) > 1 else "01_Data/raw/WA_Fn-UseC_-HR-Employee-Attrition.csv"
OUT_LKP = sys.argv[2] if len(sys.argv) > 2 else "01_Data/lookups/HR_Lookups.xlsx"
OUT_PRF = sys.argv[3] if len(sys.argv) > 3 else "02_Excel_Profiling/HR_Data_Profiling.xlsx"

# ---------------------------------------------------------------- styles
FONT = "Arial"
NAVY = "1F3864"
F_BASE = Font(name=FONT, size=10)
F_BOLD = Font(name=FONT, size=10, bold=True)
F_HDR = Font(name=FONT, size=10, bold=True, color="FFFFFF")
F_TITLE = Font(name=FONT, size=14, bold=True, color=NAVY)
F_SUB = Font(name=FONT, size=11, bold=True, color=NAVY)
F_NOTE = Font(name=FONT, size=9, italic=True, color="595959")
F_INPUT = Font(name=FONT, size=10, color="0000FF")          # hardcoded inputs (band bounds, codes)
FILL_HDR = PatternFill("solid", fgColor=NAVY)
FILL_TOT = PatternFill("solid", fgColor="D9E1F2")
FILL_PASS = PatternFill("solid", fgColor="C6EFCE")
FILL_FAIL = PatternFill("solid", fgColor="FFC7CE")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)
PCT = "0.0%"
PTS = '+0.0%;-0.0%;0.0%'
INT = "#,##0"
DEC = "0.00"

def style_header(ws, row, c1, c2):
    for c in range(c1, c2 + 1):
        cell = ws.cell(row=row, column=c)
        cell.font, cell.fill, cell.border = F_HDR, FILL_HDR, BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

def style_body(ws, r1, r2, c1, c2):
    for r in range(r1, r2 + 1):
        for c in range(c1, c2 + 1):
            cell = ws.cell(row=r, column=c)
            is_input = cell.font is not None and cell.font.color is not None and cell.font.color.rgb == "000000FF"
            if not is_input:
                cell.font = F_BASE
            cell.border = BORDER

def set_widths(ws, widths):
    for col, w in widths.items():
        ws.column_dimensions[col].width = w

def add_table(ws, name, ref, style="TableStyleMedium2"):
    t = Table(displayName=name, ref=ref)
    t.tableStyleInfo = TableStyleInfo(name=style, showRowStripes=True)
    ws.add_table(t)

# ---------------------------------------------------------------- lookup definitions
# Code labels documented on the dataset's Kaggle page (IBM HR Analytics Employee Attrition & Performance).
EDUCATION = [(1, "Below College"), (2, "College"), (3, "Bachelor"), (4, "Master"), (5, "Doctor")]
SATISFACTION = [(1, "Low"), (2, "Medium"), (3, "High"), (4, "Very High")]
WLB = [(1, "Bad"), (2, "Good"), (3, "Better"), (4, "Best")]
PERFORMANCE = [(1, "Low"), (2, "Good"), (3, "Excellent"), (4, "Outstanding")]
TRAVEL = [("Non-Travel", "No Travel"), ("Travel_Rarely", "Travels Rarely"),
          ("Travel_Frequently", "Travels Frequently")]
# Band definitions (project design choices, inclusive integer bounds)
AGE_BANDS = [("18-25", 18, 25), ("26-35", 26, 35), ("36-45", 36, 45), ("46-55", 46, 55), ("56-60", 56, 60)]
TENURE_BANDS = [("0-1 yrs", 0, 1), ("2-5 yrs", 2, 5), ("6-10 yrs", 6, 10), ("11+ yrs", 11, 99)]
INCOME_BANDS = [("Under 3,000", 0, 2999), ("3,000-5,999", 3000, 5999), ("6,000-9,999", 6000, 9999),
                ("10,000-14,999", 10000, 14999), ("15,000+", 15000, 999999)]
TRAINING_BANDS = [("None (0)", 0, 0), ("Low (1-2)", 1, 2), ("Medium (3-4)", 3, 4), ("High (5-6)", 5, 6)]

# =====================================================================================
# 1) HR_Lookups.xlsx
# =====================================================================================
def build_lookups():
    wb = Workbook()
    rd = wb.active
    rd.title = "README"
    rd["A1"] = "HR Lookups - Talent & Learning Analytics Dashboard"; rd["A1"].font = F_TITLE
    notes = [
        "Purpose: mapping tables loaded by Power Query (Step 3) to turn numeric codes into readable labels and to assign bands.",
        "Each sheet holds one Excel table. Power Query should load the TABLES (names below), not the sheets.",
        "Source of code labels: dataset documentation on Kaggle - 'IBM HR Analytics Employee Attrition & Performance'.",
        "Band tables (Age, Tenure, Income, Training) are project design choices. Bounds are inclusive integers: Min <= value <= Max.",
        "Blue text = hardcoded input values. If you change a band, re-run Step 2 validation so the profiling workbook stays in sync.",
        "MonthlyIncome currency is not stated in the dataset, so income bands carry no currency symbol.",
        "SortOrder columns are used in Power BI 'Sort by column' so bands appear in logical order (not alphabetical).",
    ]
    for i, n in enumerate(notes, start=3):
        rd.cell(row=i, column=1, value=n).font = F_BASE
    rd.cell(row=11, column=1, value="Table").font = F_BOLD
    rd.cell(row=11, column=2, value="Sheet").font = F_BOLD
    rd.cell(row=11, column=3, value="Applies to column(s) in raw data").font = F_BOLD
    rd.cell(row=11, column=4, value="Rows").font = F_BOLD
    index = [
        ("Lkp_Education", "Education", "Education", len(EDUCATION)),
        ("Lkp_Satisfaction", "Satisfaction", "EnvironmentSatisfaction, JobSatisfaction, RelationshipSatisfaction, JobInvolvement", len(SATISFACTION)),
        ("Lkp_WorkLifeBalance", "WorkLifeBalance", "WorkLifeBalance", len(WLB)),
        ("Lkp_PerformanceRating", "PerformanceRating", "PerformanceRating", len(PERFORMANCE)),
        ("Lkp_BusinessTravel", "BusinessTravel", "BusinessTravel", len(TRAVEL)),
        ("Lkp_AgeBand", "AgeBand", "Age", len(AGE_BANDS)),
        ("Lkp_TenureBand", "TenureBand", "YearsAtCompany", len(TENURE_BANDS)),
        ("Lkp_IncomeBand", "IncomeBand", "MonthlyIncome", len(INCOME_BANDS)),
        ("Lkp_TrainingBand", "TrainingBand", "TrainingTimesLastYear", len(TRAINING_BANDS)),
    ]
    for i, row in enumerate(index, start=12):
        for j, v in enumerate(row, start=1):
            rd.cell(row=i, column=j, value=v).font = F_BASE
    set_widths(rd, {"A": 26, "B": 20, "C": 80, "D": 8})

    def code_sheet(title, table, headers, rows):
        ws = wb.create_sheet(title)
        ws.append(headers)
        for k, r in enumerate(rows, start=1):
            ws.append(list(r) + [k])
        n = len(rows) + 1
        style_header(ws, 1, 1, len(headers))
        style_body(ws, 2, n, 1, len(headers))
        for r in range(2, n + 1):
            for c in range(1, len(headers) + 1):
                ws.cell(row=r, column=c).font = F_INPUT
        add_table(ws, table, f"A1:{get_column_letter(len(headers))}{n}")
        set_widths(ws, {get_column_letter(i): 22 for i in range(1, len(headers) + 1)})
        ws.freeze_panes = "A2"

    code_sheet("Education", "Lkp_Education", ["EducationCode", "EducationLevel", "SortOrder"], EDUCATION)
    code_sheet("Satisfaction", "Lkp_Satisfaction", ["RatingCode", "RatingLabel", "SortOrder"], SATISFACTION)
    code_sheet("WorkLifeBalance", "Lkp_WorkLifeBalance", ["WLBCode", "WLBLabel", "SortOrder"], WLB)
    code_sheet("PerformanceRating", "Lkp_PerformanceRating", ["PerformanceCode", "PerformanceLabel", "SortOrder"], PERFORMANCE)
    code_sheet("BusinessTravel", "Lkp_BusinessTravel", ["BusinessTravelCode", "BusinessTravelLabel", "SortOrder"], TRAVEL)
    code_sheet("AgeBand", "Lkp_AgeBand", ["AgeBand", "MinAge", "MaxAge", "SortOrder"], AGE_BANDS)
    code_sheet("TenureBand", "Lkp_TenureBand", ["TenureBand", "MinYears", "MaxYears", "SortOrder"], TENURE_BANDS)
    code_sheet("IncomeBand", "Lkp_IncomeBand", ["IncomeBand", "MinIncome", "MaxIncome", "SortOrder"], INCOME_BANDS)
    code_sheet("TrainingBand", "Lkp_TrainingBand", ["TrainingBand", "MinTrainings", "MaxTrainings", "SortOrder"], TRAINING_BANDS)
    wb.save(OUT_LKP)

# =====================================================================================
# 2) HR_Data_Profiling.xlsx
# =====================================================================================
DESCRIPTIONS = {
    "Age": "Employee age in years",
    "Attrition": "Did the employee leave? (Yes/No) - target variable",
    "BusinessTravel": "Travel frequency for work",
    "DailyRate": "Daily rate (meaning undocumented in source)",
    "Department": "Department",
    "DistanceFromHome": "Distance from home to work (units undocumented)",
    "Education": "Education level code 1-5 (see Lkp_Education)",
    "EducationField": "Field of education",
    "EmployeeCount": "Always 1 (constant)",
    "EmployeeNumber": "Unique employee identifier",
    "EnvironmentSatisfaction": "Work-environment satisfaction 1-4",
    "Gender": "Gender",
    "HourlyRate": "Hourly rate (meaning undocumented in source)",
    "JobInvolvement": "Job involvement 1-4",
    "JobLevel": "Job level 1-5 (1 = most junior)",
    "JobRole": "Job role",
    "JobSatisfaction": "Job satisfaction 1-4",
    "MaritalStatus": "Marital status",
    "MonthlyIncome": "Monthly income (currency undocumented)",
    "MonthlyRate": "Monthly rate (meaning undocumented in source)",
    "NumCompaniesWorked": "Number of previous employers",
    "Over18": "Always 'Y' (constant)",
    "OverTime": "Works overtime? (Yes/No)",
    "PercentSalaryHike": "Last salary increase in %",
    "PerformanceRating": "Performance rating 1-4 (only 3 and 4 occur)",
    "RelationshipSatisfaction": "Relationship satisfaction 1-4",
    "StandardHours": "Always 80 (constant)",
    "StockOptionLevel": "Stock option level 0-3",
    "TotalWorkingYears": "Total years of work experience",
    "TrainingTimesLastYear": "Number of trainings attended last year (0-6) - the learning metric",
    "WorkLifeBalance": "Work-life balance 1-4",
    "YearsAtCompany": "Tenure at this company in years",
    "YearsInCurrentRole": "Years in current role",
    "YearsSinceLastPromotion": "Years since last promotion",
    "YearsWithCurrManager": "Years with current manager",
}
PQ_ACTION = {
    "Age": "Keep + band -> AgeBand (Lkp_AgeBand)",
    "Attrition": "Keep + add AttritionFlag (Yes=1, No=0)",
    "BusinessTravel": "Map label via Lkp_BusinessTravel",
    "DailyRate": "Drop - undocumented, not used",
    "Department": "Move to DimDepartment",
    "DistanceFromHome": "Keep (fact)",
    "Education": "Map label via Lkp_Education -> DimEducation",
    "EducationField": "Keep (DimEducation attribute)",
    "EmployeeCount": "Drop - constant",
    "EmployeeNumber": "Keep, rename EmployeeID (primary key)",
    "EnvironmentSatisfaction": "Map label via Lkp_Satisfaction",
    "Gender": "Keep (DimEmployee attribute)",
    "HourlyRate": "Drop - undocumented, not used",
    "JobInvolvement": "Map label via Lkp_Satisfaction",
    "JobLevel": "Keep (fact / slicer)",
    "JobRole": "Move to DimJobRole",
    "JobSatisfaction": "Map label via Lkp_Satisfaction -> DimSatisfaction",
    "MaritalStatus": "Keep (DimEmployee attribute)",
    "MonthlyIncome": "Keep + band -> IncomeBand (Lkp_IncomeBand)",
    "MonthlyRate": "Drop - undocumented, not used",
    "NumCompaniesWorked": "Keep (fact)",
    "Over18": "Drop - constant",
    "OverTime": "Keep + add OverTimeFlag (Yes=1, No=0)",
    "PercentSalaryHike": "Keep (fact)",
    "PerformanceRating": "Map label via Lkp_PerformanceRating",
    "RelationshipSatisfaction": "Map label via Lkp_Satisfaction",
    "StandardHours": "Drop - constant",
    "StockOptionLevel": "Keep (fact / slicer)",
    "TotalWorkingYears": "Keep (fact)",
    "TrainingTimesLastYear": "Keep + band -> TrainingBand (Lkp_TrainingBand)",
    "WorkLifeBalance": "Map label via Lkp_WorkLifeBalance",
    "YearsAtCompany": "Keep + band -> TenureBand (Lkp_TenureBand)",
    "YearsInCurrentRole": "Keep (fact)",
    "YearsSinceLastPromotion": "Keep (fact)",
    "YearsWithCurrManager": "Keep (fact)",
}

def build_profiling(df):
    wb = Workbook()
    cols = list(df.columns)
    n = len(df)
    T = lambda c: f"tbl_HR[{c}]"          # structured reference helper
    cells = {}                              # named cells for cross-sheet references / validation

    # ---------------- README
    rd = wb.active
    rd.title = "README"
    rd["A1"] = "HR Data Profiling - Talent & Learning Analytics Dashboard"; rd["A1"].font = F_TITLE
    lines = [
        ("Purpose", "Step 2 of the project: understand the raw HR data before building the Power BI model."),
        ("Source", "IBM HR Analytics Employee Attrition & Performance (fictional dataset created by IBM data scientists; published on Kaggle). File: 01_Data/raw/WA_Fn-UseC_-HR-Employee-Attrition.csv"),
        ("Data", f"{n:,} employees x {len(cols)} columns. Point-in-time snapshot - no dates, so no trends over time."),
        ("How it works", "Raw_Data holds the CSV as Excel table 'tbl_HR'. Every number on the other sheets is a formula (COUNTIFS, AVERAGEIFS, SUMPRODUCT) over that table, so it recalculates if the data changes."),
        ("Band bounds", "Blue numbers are hardcoded band bounds / codes. They mirror 01_Data/lookups/HR_Lookups.xlsx - change both together."),
        ("Reconciliation", "KPI_Summary values are the expected results for the Power BI DAX measures in Step 5. Power BI must match them exactly."),
        ("Pivot tables", "Practice tip: Insert > PivotTable from tbl_HR (e.g. Department in Rows, Attrition in Columns, Count of EmployeeNumber) and compare with Attrition_Profile - they must match."),
        ("", ""),
        ("Sheet", "Contents"),
        ("KPI_Summary", "Headline KPIs: headcount, attrition, tenure, training, overtime, satisfaction"),
        ("Attrition_Profile", "Attrition rate by department, role, overtime, demographics, tenure, income, job level, stock options"),
        ("Training_Profile", "Training (learning) activity: attrition by training level, training by department and role"),
        ("Satisfaction_Profile", "Attrition by job / environment / relationship satisfaction, involvement, work-life balance, performance"),
        ("Data_Dictionary", "Every column: type, description, min, max, distinct values, blanks, and the planned Power Query action"),
        ("Data_Quality", "Automated PASS/FAIL data quality checks"),
        ("Key_Findings", "First observations (formula-driven text) - to be confirmed in Power BI"),
        ("Raw_Data", "Original data, unchanged (table tbl_HR)"),
    ]
    for i, (k, v) in enumerate(lines, start=3):
        rd.cell(row=i, column=1, value=k).font = F_BOLD
        rd.cell(row=i, column=2, value=v).font = F_BASE
        rd.cell(row=i, column=2).alignment = Alignment(wrap_text=True, vertical="top")
    set_widths(rd, {"A": 22, "B": 110})

    # ---------------- Raw_Data (created now so it exists for references; filled below)
    ws_raw = wb.create_sheet("Raw_Data")
    ws_raw.append(cols)
    for row in df.itertuples(index=False):
        ws_raw.append(list(row))
    last_col = get_column_letter(len(cols))
    add_table(ws_raw, "tbl_HR", f"A1:{last_col}{n + 1}", "TableStyleLight9")
    for c in range(1, len(cols) + 1):
        ws_raw.cell(row=1, column=c).font = F_BOLD
        ws_raw.column_dimensions[get_column_letter(c)].width = max(10, min(26, len(cols[c - 1]) + 3))
    for row in ws_raw.iter_rows(min_row=2, max_row=n + 1):
        for cell in row:
            cell.font = F_BASE
    ws_raw.freeze_panes = "B2"

    # ---------------- KPI_Summary
    k = wb.create_sheet("KPI_Summary", 1)
    k["A1"] = "KPI Summary"; k["A1"].font = F_TITLE
    k["A2"] = "Expected values for the Power BI DAX measures (Step 5 reconciliation baseline)."; k["A2"].font = F_NOTE
    k.append([])
    hdr = ["KPI", "Value", "Excel formula logic", "Planned DAX measure", "Definition"]
    k.append(hdr); style_header(k, 4, 1, 5)
    kpis = [
        ("Headcount", f"=COUNTA({T('EmployeeNumber')})", INT, "Count of employees", "[Headcount]", "All employees in the snapshot (leavers included)"),
        ("Leavers", f'=COUNTIF({T("Attrition")},"Yes")', INT, "Count Attrition = Yes", "[Leavers]", "Employees who left"),
        ("Active Employees", "=B5-B6", INT, "Headcount - Leavers", "[Active Employees]", "Employees still employed"),
        ("Attrition Rate", "=IFERROR(B6/B5,0)", PCT, "Leavers / Headcount", "[Attrition Rate]", "Share of employees who left"),
        ("Avg Tenure (years)", f"=AVERAGE({T('YearsAtCompany')})", DEC, "Average YearsAtCompany", "[Avg Tenure]", "Mean years at the company"),
        ("Median Tenure (years)", f"=MEDIAN({T('YearsAtCompany')})", DEC, "Median YearsAtCompany", "[Median Tenure]", "Middle value of tenure (robust to long-tenure outliers)"),
        ("Avg Tenure - Leavers (years)", f'=AVERAGEIFS({T("YearsAtCompany")},{T("Attrition")},"Yes")', DEC, "Average tenure where Attrition = Yes", "[Avg Tenure Leavers]", "How long leavers stayed"),
        ("Avg Age", f"=AVERAGE({T('Age')})", DEC, "Average Age", "[Avg Age]", "Mean employee age"),
        ("Avg Monthly Income", f"=AVERAGE({T('MonthlyIncome')})", INT, "Average MonthlyIncome", "[Avg Monthly Income]", "Currency not stated in dataset"),
        ("Total Trainings Last Year", f"=SUM({T('TrainingTimesLastYear')})", INT, "Sum TrainingTimesLastYear", "[Total Trainings]", "Total training sessions attended"),
        ("Training per Employee", "=IFERROR(B14/B5,0)", DEC, "Total trainings / Headcount", "[Training per Emp]", "Average trainings per employee last year"),
        ("% Zero Training", f'=IFERROR(COUNTIF({T("TrainingTimesLastYear")},0)/B5,0)', PCT, "Employees with 0 trainings / Headcount", "[% Zero Training]", "Share of employees with no training last year"),
        ("Overtime %", f'=IFERROR(COUNTIF({T("OverTime")},"Yes")/B5,0)', PCT, "OverTime = Yes / Headcount", "[Overtime %]", "Share of employees working overtime"),
        ("Attrition Rate - Overtime", f'=IFERROR(COUNTIFS({T("OverTime")},"Yes",{T("Attrition")},"Yes")/COUNTIF({T("OverTime")},"Yes"),0)', PCT, "Leavers / Headcount where OverTime = Yes", "[Attrition Rate] filtered", "Attrition among overtime workers"),
        ("Attrition Rate - No Overtime", f'=IFERROR(COUNTIFS({T("OverTime")},"No",{T("Attrition")},"Yes")/COUNTIF({T("OverTime")},"No"),0)', PCT, "Leavers / Headcount where OverTime = No", "[Attrition Rate] filtered", "Attrition among non-overtime workers"),
        ("Avg Job Satisfaction (1-4)", f"=AVERAGE({T('JobSatisfaction')})", DEC, "Average JobSatisfaction", "[Avg Job Satisfaction]", "1 = Low, 4 = Very High"),
        ("% Low Job Satisfaction", f'=IFERROR(COUNTIF({T("JobSatisfaction")},1)/B5,0)', PCT, "JobSatisfaction = 1 / Headcount", "[% Low Job Satisfaction]", "Share rating job satisfaction 'Low'"),
        ("Avg Years Since Last Promotion", f"=AVERAGE({T('YearsSinceLastPromotion')})", DEC, "Average YearsSinceLastPromotion", "[Avg Years Since Promotion]", "Mean years since last promotion"),
        ("Avg Performance Rating", f"=AVERAGE({T('PerformanceRating')})", DEC, "Average PerformanceRating", "[Avg Performance Rating]", "Only ratings 3 (Excellent) and 4 (Outstanding) occur"),
    ]
    for i, (name, f, fmt, logic, dax, definition) in enumerate(kpis, start=5):
        k.cell(row=i, column=1, value=name)
        c = k.cell(row=i, column=2, value=f); c.number_format = fmt
        k.cell(row=i, column=3, value=logic)
        k.cell(row=i, column=4, value=dax)
        k.cell(row=i, column=5, value=definition)
        cells["KPI:" + name] = f"KPI_Summary!$B${i}"
    style_body(k, 5, 4 + len(kpis), 1, 5)
    for i in range(5, 5 + len(kpis)):
        k.cell(row=i, column=1).font = F_BOLD
    set_widths(k, {"A": 32, "B": 14, "C": 42, "D": 28, "E": 60})
    k.freeze_panes = "A5"
    HC = cells["KPI:Headcount"]
    CO_RATE = cells["KPI:Attrition Rate"]

    # ---------------- generic block writer for profile sheets
    # columns: A Category | B Code/Min | C Max | D Headcount | E Leavers | F Attrition Rate | G vs Company | H % of Headcount
    PROFILE_HDR = ["Category", "Code / Min", "Max", "Headcount", "Leavers", "Attrition Rate", "vs Company (pts)", "% of Headcount"]

    def block(ws, start_row, title, field, items, kind, note=None, show_code=True):
        """kind: 'text' (match category text), 'code' (match numeric code), 'band' (Min<=x<=Max)
        items: list of label (text) | (code,label) | (label,min,max)."""
        ws.cell(row=start_row, column=1, value=title).font = F_SUB
        r = start_row + 1
        if note:
            ws.cell(row=r, column=1, value=note).font = F_NOTE
            r += 1
        for j, h in enumerate(PROFILE_HDR, start=1):
            ws.cell(row=r, column=j, value=h)
        style_header(ws, r, 1, 8)
        first = r + 1
        rows_meta = []
        for it in items:
            r += 1
            if kind == "text":
                ws.cell(row=r, column=1, value=it)
                crit = f"{T(field)},$A{r}"
            elif kind == "code":
                code, label = it
                disp = f"{code} - {label}" if show_code else label
                ws.cell(row=r, column=1, value=disp)
                ws.cell(row=r, column=2, value=code).font = F_INPUT
                crit = f"{T(field)},$B{r}"
            else:
                label, lo, hi = it
                ws.cell(row=r, column=1, value=label)
                ws.cell(row=r, column=2, value=lo).font = F_INPUT
                ws.cell(row=r, column=3, value=hi).font = F_INPUT
                crit = f'{T(field)},">="&$B{r},{T(field)},"<="&$C{r}'
            ws.cell(row=r, column=4, value=f"=COUNTIFS({crit})").number_format = INT
            ws.cell(row=r, column=5, value=f'=COUNTIFS({crit},{T("Attrition")},"Yes")').number_format = INT
            ws.cell(row=r, column=6, value=f'=IF(D{r}=0,"n/a",E{r}/D{r})').number_format = PCT
            ws.cell(row=r, column=7, value=f'=IF(D{r}=0,"n/a",F{r}-{CO_RATE})').number_format = PTS
            ws.cell(row=r, column=8, value=f"=IFERROR(D{r}/{HC},0)").number_format = PCT
            rows_meta.append((r, ws.cell(row=r, column=1).value))
        last = r
        r += 1
        ws.cell(row=r, column=1, value="Total")
        ws.cell(row=r, column=4, value=f"=SUM(D{first}:D{last})").number_format = INT
        ws.cell(row=r, column=5, value=f"=SUM(E{first}:E{last})").number_format = INT
        ws.cell(row=r, column=6, value=f'=IF(D{r}=0,"n/a",E{r}/D{r})').number_format = PCT
        ws.cell(row=r, column=7, value=f'=IF(D{r}=0,"n/a",F{r}-{CO_RATE})').number_format = PTS
        ws.cell(row=r, column=8, value=f"=IFERROR(D{r}/{HC},0)").number_format = PCT
        style_body(ws, first, r, 1, 8)
        for c in range(1, 9):
            ws.cell(row=r, column=c).fill = FILL_TOT
            ws.cell(row=r, column=c).font = F_BOLD
        # colour scale on attrition rate (green low -> red high)
        ws.conditional_formatting.add(f"F{first}:F{last}", ColorScaleRule(
            start_type="min", start_color="63BE7B", mid_type="percentile", mid_value=50,
            mid_color="FFEB84", end_type="max", end_color="F8696B"))
        return {"title": title, "field": field, "kind": kind, "first": first, "last": last,
                "total": r, "rows": rows_meta, "next": r + 3}

    def bar_chart(ws, blk, anchor, title, height=7.0, width=14.0):
        ch = BarChart(); ch.type = "bar"; ch.style = 10
        ch.title = title; ch.y_axis.title = "Attrition rate"; ch.y_axis.numFmt = "0%"
        ch.y_axis.majorGridlines = None
        data = Reference(ws, min_col=6, min_row=blk["first"] - 1, max_row=blk["last"])
        cats = Reference(ws, min_col=1, min_row=blk["first"], max_row=blk["last"])
        ch.add_data(data, titles_from_data=True); ch.set_categories(cats)
        ch.legend = None; ch.height, ch.width = height, width
        ws.add_chart(ch, anchor)

    def sorted_values(col):
        return sorted(df[col].unique().tolist())

    blocks = {}

    # ---------------- Attrition_Profile
    a = wb.create_sheet("Attrition_Profile", 2)
    a["A1"] = "Attrition Profile"; a["A1"].font = F_TITLE
    a["A2"] = "Attrition rate = Leavers / Headcount. 'vs Company' = segment rate minus company rate, in percentage points. Colour: green = low, red = high attrition."
    a["A2"].font = F_NOTE
    r = 4
    spec = [
        ("Department", "Department", "text", sorted_values("Department"), None),
        ("Job Role", "JobRole", "text", sorted_values("JobRole"), None),
        ("Overtime", "OverTime", "text", ["No", "Yes"], None),
        ("Business Travel", "BusinessTravel", "text", [c for c, _ in TRAVEL], None),
        ("Gender", "Gender", "text", sorted_values("Gender"), None),
        ("Marital Status", "MaritalStatus", "text", sorted_values("MaritalStatus"), None),
        ("Age Band", "Age", "band", AGE_BANDS, "Bands mirror Lkp_AgeBand"),
        ("Tenure Band (YearsAtCompany)", "YearsAtCompany", "band", TENURE_BANDS, "Bands mirror Lkp_TenureBand"),
        ("Monthly Income Band", "MonthlyIncome", "band", INCOME_BANDS, "Bands mirror Lkp_IncomeBand"),
        ("Job Level", "JobLevel", "code", [(v, f"Level {v}") for v in range(1, 6)], "1 = most junior, 5 = most senior"),
        ("Stock Option Level", "StockOptionLevel", "code", [(v, f"Level {v}") for v in range(0, 4)], "0 = no stock options"),
    ]
    for title, field, kind, items, note in spec:
        b = block(a, r, title, field, items, kind, note, show_code=field not in ("JobLevel", "StockOptionLevel"))
        blocks["Attrition:" + title] = b
        r = b["next"]
    bar_chart(a, blocks["Attrition:Department"], "J4", "Attrition rate by Department")
    bar_chart(a, blocks["Attrition:Job Role"], "J20", "Attrition rate by Job Role", height=9.0)
    bar_chart(a, blocks["Attrition:Overtime"], "J40", "Attrition rate by Overtime")
    bar_chart(a, blocks["Attrition:Tenure Band (YearsAtCompany)"], "J56", "Attrition rate by Tenure")
    set_widths(a, {"A": 30, "B": 11, "C": 9, "D": 12, "E": 10, "F": 13, "G": 15, "H": 14})
    a.freeze_panes = "A4"

    # ---------------- Training_Profile
    t = wb.create_sheet("Training_Profile", 3)
    t["A1"] = "Training & Learning Profile"; t["A1"].font = F_TITLE
    t["A2"] = "Learning metric = TrainingTimesLastYear (number of trainings attended last year, 0-6). It is the only learning field in the dataset."
    t["A2"].font = F_NOTE
    b1 = block(t, 4, "Attrition by Trainings Attended Last Year", "TrainingTimesLastYear",
               [(v, f"{v} training" + ("" if v == 1 else "s")) for v in range(0, 7)], "code", show_code=False)
    blocks["Training:ByCount"] = b1
    b2 = block(t, b1["next"], "Attrition by Training Band", "TrainingTimesLastYear", TRAINING_BANDS, "band",
               "Bands mirror Lkp_TrainingBand")
    blocks["Training:ByBand"] = b2
    bar_chart(t, b2, "J4", "Attrition rate by Training Band")

    def training_by(ws, start_row, title, field, values):
        ws.cell(row=start_row, column=1, value=title).font = F_SUB
        hr = start_row + 1
        hdrs = [title.split(" by ")[-1], "Headcount", "Total Trainings", "Training per Employee",
                "% Zero Training", "% High Training (5-6)", "Attrition Rate"]
        for j, h in enumerate(hdrs, start=1):
            ws.cell(row=hr, column=j, value=h)
        style_header(ws, hr, 1, 7)
        rr = hr
        for v in values:
            rr += 1
            ws.cell(row=rr, column=1, value=v)
            ws.cell(row=rr, column=2, value=f"=COUNTIFS({T(field)},$A{rr})").number_format = INT
            ws.cell(row=rr, column=3, value=f"=SUMIFS({T('TrainingTimesLastYear')},{T(field)},$A{rr})").number_format = INT
            ws.cell(row=rr, column=4, value=f"=IFERROR(C{rr}/B{rr},0)").number_format = DEC
            ws.cell(row=rr, column=5, value=f"=IFERROR(COUNTIFS({T(field)},$A{rr},{T('TrainingTimesLastYear')},0)/B{rr},0)").number_format = PCT
            ws.cell(row=rr, column=6, value=f'=IFERROR(COUNTIFS({T(field)},$A{rr},{T("TrainingTimesLastYear")},">=5")/B{rr},0)').number_format = PCT
            ws.cell(row=rr, column=7, value=f'=IFERROR(COUNTIFS({T(field)},$A{rr},{T("Attrition")},"Yes")/B{rr},0)').number_format = PCT
        first, last = hr + 1, rr
        rr += 1
        ws.cell(row=rr, column=1, value="Total")
        ws.cell(row=rr, column=2, value=f"=SUM(B{first}:B{last})").number_format = INT
        ws.cell(row=rr, column=3, value=f"=SUM(C{first}:C{last})").number_format = INT
        ws.cell(row=rr, column=4, value=f"=IFERROR(C{rr}/B{rr},0)").number_format = DEC
        ws.cell(row=rr, column=5, value=f"=IFERROR(COUNTIF({T('TrainingTimesLastYear')},0)/B{rr},0)").number_format = PCT
        ws.cell(row=rr, column=6, value=f'=IFERROR(COUNTIF({T("TrainingTimesLastYear")},">=5")/B{rr},0)').number_format = PCT
        ws.cell(row=rr, column=7, value=f'=IFERROR(COUNTIF({T("Attrition")},"Yes")/B{rr},0)').number_format = PCT
        style_body(ws, first, rr, 1, 7)
        for c in range(1, 8):
            ws.cell(row=rr, column=c).fill = FILL_TOT; ws.cell(row=rr, column=c).font = F_BOLD
        ws.conditional_formatting.add(f"D{first}:D{last}", ColorScaleRule(
            start_type="min", start_color="F8696B", mid_type="percentile", mid_value=50,
            mid_color="FFEB84", end_type="max", end_color="63BE7B"))
        return {"first": first, "last": last, "total": rr, "next": rr + 3, "values": values, "field": field}

    b3 = training_by(t, b2["next"], "Training by Department", "Department", sorted_values("Department"))
    b4 = training_by(t, b3["next"], "Training by Job Role", "JobRole", sorted_values("JobRole"))
    blocks["Training:ByDept"], blocks["Training:ByRole"] = b3, b4
    set_widths(t, {"A": 30, "B": 12, "C": 14, "D": 14, "E": 12, "F": 15, "G": 15, "H": 14})
    t.freeze_panes = "A4"

    # ---------------- Satisfaction_Profile
    s = wb.create_sheet("Satisfaction_Profile", 4)
    s["A1"] = "Satisfaction & Performance Profile"; s["A1"].font = F_TITLE
    s["A2"] = "Codes mapped with labels from HR_Lookups.xlsx. 'n/a' = code defined but no employees have it."
    s["A2"].font = F_NOTE
    r = 4
    sspec = [
        ("Job Satisfaction", "JobSatisfaction", SATISFACTION),
        ("Environment Satisfaction", "EnvironmentSatisfaction", SATISFACTION),
        ("Relationship Satisfaction", "RelationshipSatisfaction", SATISFACTION),
        ("Job Involvement", "JobInvolvement", SATISFACTION),
        ("Work-Life Balance", "WorkLifeBalance", WLB),
        ("Performance Rating", "PerformanceRating", PERFORMANCE),
        ("Education", "Education", EDUCATION),
    ]
    for title, field, items in sspec:
        b = block(s, r, title, field, items, "code")
        blocks["Satisfaction:" + title] = b
        r = b["next"]
    bar_chart(s, blocks["Satisfaction:Job Satisfaction"], "J4", "Attrition rate by Job Satisfaction")
    bar_chart(s, blocks["Satisfaction:Work-Life Balance"], "J20", "Attrition rate by Work-Life Balance")
    set_widths(s, {"A": 30, "B": 11, "C": 9, "D": 12, "E": 10, "F": 13, "G": 15, "H": 14})
    s.freeze_panes = "A4"

    # ---------------- Data_Dictionary
    dd = wb.create_sheet("Data_Dictionary", 5)
    dd["A1"] = "Data Dictionary"; dd["A1"].font = F_TITLE
    dd["A2"] = "Min / Max / Distinct / Blanks are live formulas over tbl_HR. 'Planned Power Query action' is the Step 3 design."
    dd["A2"].font = F_NOTE
    h = ["#", "Column", "Data type", "Description", "Min", "Max", "Distinct values", "Blanks", "Planned Power Query action"]
    for j, v in enumerate(h, start=1):
        dd.cell(row=4, column=j, value=v)
    style_header(dd, 4, 1, len(h))
    for i, c in enumerate(cols, start=1):
        r = 4 + i
        numeric = pd.api.types.is_numeric_dtype(df[c])
        dd.cell(row=r, column=1, value=i)
        dd.cell(row=r, column=2, value=c)
        dd.cell(row=r, column=3, value="Whole number" if numeric else "Text")
        dd.cell(row=r, column=4, value=DESCRIPTIONS[c])
        dd.cell(row=r, column=5, value=f"=MIN({T(c)})" if numeric else "-")
        dd.cell(row=r, column=6, value=f"=MAX({T(c)})" if numeric else "-")
        dd.cell(row=r, column=7, value=f'=SUMPRODUCT(({T(c)}<>"")/COUNTIF({T(c)},{T(c)}&""))').number_format = INT
        dd.cell(row=r, column=8, value=f"=COUNTBLANK({T(c)})").number_format = INT
        dd.cell(row=r, column=9, value=PQ_ACTION[c])
        if PQ_ACTION[c].startswith("Drop"):
            dd.cell(row=r, column=9).font = Font(name=FONT, size=10, color="C00000")
    style_body(dd, 5, 4 + len(cols), 1, len(h))
    for r in range(5, 5 + len(cols)):
        if str(dd.cell(row=r, column=9).value).startswith("Drop"):
            dd.cell(row=r, column=9).font = Font(name=FONT, size=10, color="C00000", bold=True)
    set_widths(dd, {"A": 5, "B": 26, "C": 14, "D": 58, "E": 9, "F": 9, "G": 10, "H": 9, "I": 46})
    dd.freeze_panes = "C5"
    cells["DD:first"], cells["DD:last"] = 5, 4 + len(cols)

    # ---------------- Data_Quality
    q = wb.create_sheet("Data_Quality", 6)
    q["A1"] = "Data Quality Checks"; q["A1"].font = F_TITLE
    q["A2"] = "Live checks over tbl_HR. Result = PASS when Actual equals Expected."; q["A2"].font = F_NOTE
    h = ["#", "Check", "Actual", "Expected", "Result", "Why it matters"]
    for j, v in enumerate(h, start=1):
        q.cell(row=4, column=j, value=v)
    style_header(q, 4, 1, len(h))
    total_cells = n * len(cols)
    checks = [
        ("Row count", f"=ROWS({T('EmployeeNumber')})", n, "Confirms the full file loaded"),
        ("Column count", "=COLUMNS(tbl_HR)", len(cols), "Confirms no columns lost"),
        ("Blank cells in whole table", "=COUNTBLANK(tbl_HR)", 0, "No missing values to impute"),
        ("Duplicate EmployeeNumber rows", f"=SUMPRODUCT(--(COUNTIF({T('EmployeeNumber')},{T('EmployeeNumber')})>1))", 0, "EmployeeNumber can be the primary key"),
        ("EmployeeCount = 1 for all rows", f"=COUNTIF({T('EmployeeCount')},1)", n, "Constant column -> drop"),
        ("Over18 = 'Y' for all rows", f'=COUNTIF({T("Over18")},"Y")', n, "Constant column -> drop"),
        ("StandardHours = 80 for all rows", f"=COUNTIF({T('StandardHours')},80)", n, "Constant column -> drop"),
        ("Attrition is Yes or No", f'=COUNTIF({T("Attrition")},"Yes")+COUNTIF({T("Attrition")},"No")', n, "Clean flag for attrition measures"),
        ("OverTime is Yes or No", f'=COUNTIF({T("OverTime")},"Yes")+COUNTIF({T("OverTime")},"No")', n, "Clean flag for overtime measures"),
        ("Age between 18 and 60", f'=COUNTIFS({T("Age")},">=18",{T("Age")},"<=60")', n, "All employees fall into an age band"),
        ("YearsAtCompany > TotalWorkingYears", f"=SUMPRODUCT(--({T('YearsAtCompany')}>{T('TotalWorkingYears')}))", 0, "Tenure cannot exceed career length"),
        ("YearsInCurrentRole > YearsAtCompany", f"=SUMPRODUCT(--({T('YearsInCurrentRole')}>{T('YearsAtCompany')}))", 0, "Role tenure cannot exceed company tenure"),
        ("YearsSinceLastPromotion > YearsAtCompany", f"=SUMPRODUCT(--({T('YearsSinceLastPromotion')}>{T('YearsAtCompany')}))", 0, "Promotion gap cannot exceed company tenure"),
        ("YearsWithCurrManager > YearsAtCompany", f"=SUMPRODUCT(--({T('YearsWithCurrManager')}>{T('YearsAtCompany')}))", 0, "Manager tenure cannot exceed company tenure"),
        ("Started work before age 18 (Age - TotalWorkingYears < 18)", f"=SUMPRODUCT(--(({T('Age')}-{T('TotalWorkingYears')})<18))", 0, "Plausibility of career length"),
        ("Education code outside 1-5", f'=ROWS({T("Education")})-COUNTIFS({T("Education")},">=1",{T("Education")},"<=5")', 0, "Every code maps to Lkp_Education"),
        ("Satisfaction codes outside 1-4 (4 columns)", f'=4*ROWS({T("JobSatisfaction")})-COUNTIFS({T("JobSatisfaction")},">=1",{T("JobSatisfaction")},"<=4")-COUNTIFS({T("EnvironmentSatisfaction")},">=1",{T("EnvironmentSatisfaction")},"<=4")-COUNTIFS({T("RelationshipSatisfaction")},">=1",{T("RelationshipSatisfaction")},"<=4")-COUNTIFS({T("JobInvolvement")},">=1",{T("JobInvolvement")},"<=4")', 0, "Every code maps to Lkp_Satisfaction"),
        ("TrainingTimesLastYear outside 0-6", f'=ROWS({T("TrainingTimesLastYear")})-COUNTIFS({T("TrainingTimesLastYear")},">=0",{T("TrainingTimesLastYear")},"<=6")', 0, "Every employee maps to a training band"),
        ("Attrition profile totals = Headcount", None, n, "Profile blocks cover every employee exactly once"),
    ]
    for i, (name, f, exp, why) in enumerate(checks, start=1):
        r = 4 + i
        q.cell(row=r, column=1, value=i)
        q.cell(row=r, column=2, value=name)
        if f is None:
            # every attrition block total must equal headcount -> use MIN and MAX of the totals
            refs = ",".join(f"Attrition_Profile!$D${b['total']}" for key, b in blocks.items() if key.startswith("Attrition:"))
            f = f"=IF(MIN({refs})=MAX({refs}),MIN({refs}),-1)"
        q.cell(row=r, column=3, value=f).number_format = INT
        q.cell(row=r, column=4, value=exp).number_format = INT
        q.cell(row=r, column=5, value=f'=IF(C{r}=D{r},"PASS","FAIL")')
        q.cell(row=r, column=6, value=why)
    last_q = 4 + len(checks)
    style_body(q, 5, last_q, 1, 6)
    q.conditional_formatting.add(f"E5:E{last_q}", CellIsRule(operator="equal", formula=['"PASS"'], fill=FILL_PASS))
    q.conditional_formatting.add(f"E5:E{last_q}", CellIsRule(operator="equal", formula=['"FAIL"'], fill=FILL_FAIL))
    r = last_q + 2
    q.cell(row=r, column=2, value="Checks passed").font = F_BOLD
    q.cell(row=r, column=3, value=f'=COUNTIF(E5:E{last_q},"PASS")').font = F_BOLD
    q.cell(row=r, column=4, value=len(checks)).font = F_BOLD
    q.cell(row=r, column=5, value=f'=IF(C{r}=D{r},"ALL PASS","CHECK FAILURES")').font = F_BOLD
    cells["DQ:summary_row"] = r
    set_widths(q, {"A": 5, "B": 52, "C": 12, "D": 12, "E": 10, "F": 50})
    q.freeze_panes = "A5"

    # ---------------- Key_Findings (formula-driven text)
    kf = wb.create_sheet("Key_Findings", 7)
    kf["A1"] = "Key Findings (first look)"; kf["A1"].font = F_TITLE
    kf["A2"] = "Sentences are formulas, so the numbers update with the data. Confirm every finding in Power BI before presenting it."
    kf["A2"].font = F_NOTE

    def ref(sheet, b, label, col):
        for rr, lab in b["rows"]:
            if lab == label:
                return f"{sheet}!${col}${rr}"
        raise KeyError(label)

    AP, TP, SP = "Attrition_Profile", "Training_Profile", "Satisfaction_Profile"
    ba = blocks
    rate = lambda sh, blk, lab: ref(sh, blk, lab, "F")
    head = lambda sh, blk, lab: ref(sh, blk, lab, "D")
    P = lambda x: f'TEXT({x},"0.0%")'
    dept_b, role_b = ba["Attrition:Department"], ba["Attrition:Job Role"]
    dept_rates = f"{AP}!$F${dept_b['first']}:$F${dept_b['last']}"
    dept_names = f"{AP}!$A${dept_b['first']}:$A${dept_b['last']}"
    role_rates = f"{AP}!$F${role_b['first']}:$F${role_b['last']}"
    role_names = f"{AP}!$A${role_b['first']}:$A${role_b['last']}"
    findings = [
        ("Overall attrition",
         f'="Overall attrition is "&{P(CO_RATE)}&" ("&TEXT({cells["KPI:Leavers"]},"#,##0")&" of "&TEXT({HC},"#,##0")&" employees left)."'),
        ("Overtime (driver)",
         f'="Employees working overtime leave at "&{P(rate(AP, ba["Attrition:Overtime"], "Yes"))}&" vs "&{P(rate(AP, ba["Attrition:Overtime"], "No"))}&" without overtime - about "&TEXT({rate(AP, ba["Attrition:Overtime"], "Yes")}/{rate(AP, ba["Attrition:Overtime"], "No")},"0.0")&"x higher. "&{P(cells["KPI:Overtime %"])}&" of staff work overtime."'),
        ("Early tenure (driver)",
         f'="Employees in their first 0-1 years leave at "&{P(rate(AP, ba["Attrition:Tenure Band (YearsAtCompany)"], "0-1 yrs"))}&" vs "&{P(rate(AP, ba["Attrition:Tenure Band (YearsAtCompany)"], "11+ yrs"))}&" for 11+ years. Leavers stayed "&TEXT({cells["KPI:Avg Tenure - Leavers (years)"]},"0.0")&" years on average vs "&TEXT({cells["KPI:Avg Tenure (years)"]},"0.0")&" overall."'),
        ("Low job satisfaction (driver)",
         f'="Employees with Low job satisfaction leave at "&{P(rate(SP, ba["Satisfaction:Job Satisfaction"], "1 - Low"))}&" vs "&{P(rate(SP, ba["Satisfaction:Job Satisfaction"], "4 - Very High"))}&" for Very High satisfaction."'),
        ("Department & role",
         f'="Highest-attrition department: "&INDEX({dept_names},MATCH(MAX({dept_rates}),{dept_rates},0))&" ("&{P(f"MAX({dept_rates})")}&"). Highest-attrition role: "&INDEX({role_names},MATCH(MAX({role_rates}),{role_rates},0))&" ("&{P(f"MAX({role_rates})")}&")."'),
        ("Training (learning)",
         f'="Employees with no training last year leave at "&{P(rate(TP, ba["Training:ByBand"], "None (0)"))}&" vs "&{P(rate(TP, ba["Training:ByBand"], "High (5-6)"))}&" for 5-6 trainings. Only "&{P(cells["KPI:% Zero Training"])}&" had zero training, so the no-training group is small ("&TEXT({head(TP, ba["Training:ByBand"], "None (0)")},"0")&" people) - treat as indicative."'),
        ("Junior level & stock options",
         f'="Job Level 1 attrition is "&{P(rate(AP, ba["Attrition:Job Level"], "Level 1"))}&"; employees with no stock options (level 0) leave at "&{P(rate(AP, ba["Attrition:Stock Option Level"], "Level 0"))}&"."'),
        ("Performance is NOT a driver",
         f'="Attrition is similar for Excellent ("&{P(rate(SP, ba["Satisfaction:Performance Rating"], "3 - Excellent"))}&") and Outstanding ("&{P(rate(SP, ba["Satisfaction:Performance Rating"], "4 - Outstanding"))}&") performers - high performers are not being retained better."'),
    ]
    for j, v in enumerate(["#", "Theme", "Finding"], start=1):
        kf.cell(row=4, column=j, value=v)
    style_header(kf, 4, 1, 3)
    for i, (theme, f) in enumerate(findings, start=1):
        r = 4 + i
        kf.cell(row=r, column=1, value=i)
        kf.cell(row=r, column=2, value=theme)
        kf.cell(row=r, column=3, value=f)
        for c in (1, 2, 3):
            kf.cell(row=r, column=c).alignment = Alignment(wrap_text=True, vertical="top")
        kf.row_dimensions[r].height = 30
    style_body(kf, 5, 4 + len(findings), 1, 3)
    for r in range(5, 5 + len(findings)):
        kf.cell(row=r, column=2).font = F_BOLD
    r = 6 + len(findings)
    kf.cell(row=r, column=2, value="Next").font = F_BOLD
    kf.cell(row=r, column=3, value="Recommendations are written in Step 7 after the Power BI model confirms these numbers.").font = F_NOTE
    set_widths(kf, {"A": 5, "B": 30, "C": 130})

    # sheet order: README, KPI, Attrition, Training, Satisfaction, Dictionary, Quality, Findings, Raw
    order = ["README", "KPI_Summary", "Attrition_Profile", "Training_Profile", "Satisfaction_Profile",
             "Data_Dictionary", "Data_Quality", "Key_Findings", "Raw_Data"]
    wb._sheets = [wb[nm] for nm in order]
    for ws in wb.worksheets:
        ws.sheet_view.showGridLines = ws.title == "Raw_Data"
    wb.active = 0
    wb.save(OUT_PRF)
    return blocks


if __name__ == "__main__":
    data = pd.read_csv(CSV, encoding="utf-8-sig")
    build_lookups()
    build_profiling(data)
    print("built", OUT_LKP, OUT_PRF)
