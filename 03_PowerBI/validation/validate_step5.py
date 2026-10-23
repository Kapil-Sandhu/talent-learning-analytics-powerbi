"""
Step 5 validation - DAX measures.
  A. Definition   : 33 measures; offline TMDL parse (Microsoft serializer) == spec; names unique and no clash with columns
  B. Engine       : every measure in the OPEN Power BI model has State = Ready (no syntax / semantic error) and == spec
  C. Step 2 recon : every KPI in the Step 2 Excel KPI_Summary is reproduced exactly by its measure
  D. Independent  : every measure x 16 groupings (incl. per employee, 1,470 rows) == pandas computed from the RAW CSV
  E. Edge cases   : empty selection -> BLANK (no errors), KEEPFILTERS intersection, slicer + company baseline,
                    no NaN / infinity anywhere, negative control query fails as it must
  F. Report       : report files still pass Microsoft's official JSON schemas; no Power BI dialog / warning on screen
Usage: python validate_step5.py <project_root> <offline.json> <live.json> <schema.json> <ui.json> <report.txt>
"""
import json, math, os, re, sys
import pandas as pd
from openpyxl import load_workbook

ROOT, OFF, LIVE, SCH, UI, REPORT = sys.argv[1:7]
sys.path.insert(0, os.path.join(ROOT, "03_PowerBI", "model"))
from measures_spec import MEASURES, KPI_MAP

res, out = [], []
def check(sec, name, ok, detail=""):
    res.append(bool(ok)); out.append(f"[{'PASS' if ok else 'FAIL'}] {sec}: {name}" + (f" - {detail}" if detail else ""))

off = json.load(open(OFF, encoding="utf-8-sig"))
live = json.load(open(LIVE, encoding="utf-8-sig"))
NAMES = [m["name"] for m in MEASURES]
norm = lambda s: "\n".join(l.rstrip() for l in (s or "").replace("\r\n", "\n").strip().split("\n"))

# ------------------------------------------------------------------ A. definition
S = "A. Definition"
check(S, "33 measures defined in the spec, names unique", len(NAMES) == 33 and len(set(NAMES)) == 33)
cols = {c["name"] for t in off["tables"] for c in t["columns"]}
check(S, "No measure name equals a column name anywhere in the model", not (set(NAMES) & cols), str(set(NAMES) & cols))
check(S, "Microsoft TmdlSerializer parsed the model incl. measures", off["ok"] and len(off.get("measures", [])) == 33,
      f"ok={off['ok']} measures={len(off.get('measures', []))} err={off.get('error')}")
def same_as_spec(found, label):
    by = {m["name"]: m for m in found}
    bad = []
    for m in MEASURES:
        f = by.get(m["name"])
        if not f: bad.append(f"{m['name']}: missing"); continue
        if f["table"] != "_Measures": bad.append(f"{m['name']}: table {f['table']}")
        if norm(f["expression"]) != norm("\n".join(m["expr"])): bad.append(f"{m['name']}: expression differs")
        if f["formatString"] != m["fmt"]: bad.append(f"{m['name']}: format {f['formatString']}")
        if f["displayFolder"] != m["folder"]: bad.append(f"{m['name']}: folder {f['displayFolder']}")
        if norm(f["description"]) != norm(m["desc"]): bad.append(f"{m['name']}: description differs")
    extra = set(by) - set(NAMES)
    check(label[0], f"{label[1]}: all 33 measures == spec (home table, expression, format, folder, description)",
          not bad and not extra, "; ".join(bad[:5]) + (f" extra={extra}" if extra else ""))
same_as_spec(off.get("measures", []), (S, "Offline TMDL"))
folders = {m["folder"] for m in MEASURES}
check(S, "Measures organised in 6 display folders", len(folders) == 6, str(sorted(folders)))

# ------------------------------------------------------------------ B. engine
S = "B. Engine"
check(S, "Live model reports 33 measures", len(live["measures"]) == 33, str(len(live["measures"])))
not_ready = [(m["name"], m["state"], m["errorMessage"]) for m in live["measures"] if m["state"] != "Ready" or m["errorMessage"]]
check(S, "Every measure State = Ready in Power BI (no syntax or semantic errors)", not not_ready, str(not_ready))
same_as_spec(live["measures"], (S, "Live model"))
check(S, "All DAX test queries ran (except the deliberate negative control)", not live["errors"], str(live["errors"]))

# ------------------------------------------------------------------ independent expectation from the RAW CSV
raw = pd.read_csv(os.path.join(ROOT, "01_Data/raw/WA_Fn-UseC_-HR-Employee-Attrition.csv"), encoding="utf-8-sig")
lk = load_workbook(os.path.join(ROOT, "01_Data/lookups/HR_Lookups.xlsx"), data_only=True)
def ltab(name):
    for ws in lk.worksheets:
        if name in ws.tables:
            rows = [[c.value for c in r] for r in ws[ws.tables[name].ref]]
            return pd.DataFrame(rows[1:], columns=rows[0])
def band(series, t):
    b = ltab(t); out = []
    for v in series:
        m = b[(b.iloc[:, 1] <= v) & (v <= b.iloc[:, 2])]
        assert len(m) == 1, (t, v)
        out.append(m.iloc[0, 0])
    return out
r = raw.copy()
r["att"] = (r.Attrition == "Yes").astype(int); r["ot"] = (r.OverTime == "Yes").astype(int)
r["EducationLevel"] = r.Education.map(dict(ltab("Lkp_Education").iloc[:, :2].values))
r["JobSatLabel"] = r.JobSatisfaction.map(dict(ltab("Lkp_Satisfaction").iloc[:, :2].values))
r["WLBLabel"] = r.WorkLifeBalance.map(dict(ltab("Lkp_WorkLifeBalance").iloc[:, :2].values))
r["TravelLabel"] = r.BusinessTravel.map(dict(ltab("Lkp_BusinessTravel").iloc[:, :2].values))
r["AgeBand"] = band(r.Age, "Lkp_AgeBand"); r["TenureBand"] = band(r.YearsAtCompany, "Lkp_TenureBand")
r["IncomeBand"] = band(r.MonthlyIncome, "Lkp_IncomeBand"); r["TrainingBand"] = band(r.TrainingTimesLastYear, "Lkp_TrainingBand")
r["Status"] = r.att.map({1: "Left", 0: "Active"})
COMPANY = r.att.mean()

def rate(g):
    return (g.att.sum() / len(g)) if len(g) else None
def mean(s):
    return float(s.mean()) if len(s) else None
def divide(a, b):
    if a is None or b is None or b == 0: return None
    return a / b
def expected(g):
    hc = len(g)
    if hc == 0:
        return {n: (COMPANY if n == "Company Attrition Rate" else None) for n in NAMES}
    L = int(g.att.sum())
    e = {"Headcount": hc, "Leavers": L, "Active Employees": hc - L, "Attrition Rate": L / hc, "Retention Rate": (hc - L) / hc,
         "Company Attrition Rate": COMPANY, "Attrition vs Company (pts)": L / hc - COMPANY,
         "Avg Tenure (Years)": mean(g.YearsAtCompany), "Median Tenure (Years)": float(g.YearsAtCompany.median()),
         "Avg Tenure of Leavers (Years)": mean(g[g.att == 1].YearsAtCompany),
         "Early Tenure Attrition Rate (0-1 yrs)": rate(g[g.YearsAtCompany <= 1]),
         "Avg Years Since Promotion": mean(g.YearsSinceLastPromotion),
         "Total Trainings": int(g.TrainingTimesLastYear.sum()), "Training per Employee": g.TrainingTimesLastYear.sum() / hc,
         "Employees with No Training": int((g.TrainingTimesLastYear == 0).sum()),
         "% No Training": (g.TrainingTimesLastYear == 0).sum() / hc, "% High Training (5-6)": (g.TrainingTimesLastYear >= 5).sum() / hc,
         "Attrition Rate - No Training": rate(g[g.TrainingTimesLastYear == 0]),
         "Overtime Employees": int(g.ot.sum()), "Overtime %": g.ot.sum() / hc,
         "Attrition Rate - Overtime": rate(g[g.ot == 1]), "Attrition Rate - No Overtime": rate(g[g.ot == 0]),
         "Avg Job Satisfaction": mean(g.JobSatisfaction), "% Low Job Satisfaction": (g.JobSatisfaction == 1).sum() / hc,
         "Attrition Rate - Low Job Satisfaction": rate(g[g.JobSatisfaction == 1]),
         "Avg Environment Satisfaction": mean(g.EnvironmentSatisfaction), "Avg Work-Life Balance": mean(g.WorkLifeBalance),
         "Avg Job Involvement": mean(g.JobInvolvement), "Avg Performance Rating": mean(g.PerformanceRating),
         "Avg Age": mean(g.Age), "Avg Monthly Income": mean(g.MonthlyIncome), "Avg Salary Hike (%)": mean(g.PercentSalaryHike)}
    e["Overtime Attrition Ratio"] = divide(e["Attrition Rate - Overtime"], e["Attrition Rate - No Overtime"])
    assert set(e) == set(NAMES), set(NAMES) ^ set(e)
    return e
def eq(a, b):
    if b is None: return a is None
    if a is None: return False
    if isinstance(a, float) and (math.isnan(a) or math.isinf(a)): return False
    return abs(float(a) - float(b)) <= 1e-9 * max(1.0, abs(float(b)))

def qframe(key):
    q = live["queries"][key]
    return [dict(zip(q["columns"], row)) for row in q["rows"]]
def mval(row, n):
    return row.get(f"[{n}]")

# ------------------------------------------------------------------ C. Step 2 reconciliation
S = "C. Step 2 KPIs"
kws = load_workbook(os.path.join(ROOT, "02_Excel_Profiling/HR_Data_Profiling.xlsx"), data_only=True)["KPI_Summary"]
kpi = {kws.cell(i, 1).value: kws.cell(i, 2).value for i in range(5, 40) if kws.cell(i, 1).value}
total = qframe("total")[0]
check(S, f"All {len(KPI_MAP)} Step 2 KPIs are mapped to a measure", set(KPI_MAP) <= set(kpi) and set(KPI_MAP.values()) <= set(NAMES))
for k, mname in KPI_MAP.items():
    check(S, f"{k} (Excel {kpi[k]:.6g}) == [{mname}] in Power BI", eq(mval(total, mname), kpi[k]), f"powerbi={mval(total, mname)}")

# ------------------------------------------------------------------ D. independent comparison, every measure x grouping
S = "D. Measures vs raw CSV"
exp_total = expected(r)
bad = [n for n in NAMES if not eq(mval(total, n), exp_total[n])]
check(S, "Total level: all 33 measures == pandas", not bad, str([(n, mval(total, n), exp_total[n]) for n in bad]))
G = {"dept": ("DimDepartment[Department]", "Department"), "role": ("DimJobRole[JobRole]", "JobRole"),
     "edu": ("DimEducation[EducationLevel]", "EducationLevel"), "jobsat": ("DimJobSatisfaction[JobSatisfaction]", "JobSatLabel"),
     "age": ("DimAgeBand[AgeBand]", "AgeBand"), "tenure": ("DimTenureBand[TenureBand]", "TenureBand"),
     "income": ("DimIncomeBand[IncomeBand]", "IncomeBand"), "training": ("DimTrainingBand[TrainingBand]", "TrainingBand"),
     "overtime": ("DimEmployee[OverTime]", "OverTime"), "gender": ("DimEmployee[Gender]", "Gender"),
     "marital": ("DimEmployee[MaritalStatus]", "MaritalStatus"), "travel": ("DimEmployee[BusinessTravel]", "TravelLabel"),
     "wlb": ("DimEmployee[WorkLifeBalance]", "WLBLabel"), "status": ("DimEmployee[EmployeeStatus]", "Status"),
     "employee": ("DimEmployee[EmployeeID]", "EmployeeNumber")}
cells_checked = 0
for key, (dcol, pcol) in G.items():
    rows = qframe(f"by_{key}")
    members = sorted(r[pcol].unique().tolist())
    got_members = sorted(row[dcol] for row in rows)
    problems = [] if got_members == members else [f"members {got_members[:5]} vs {members[:5]}"]
    for row in rows:
        e = expected(r[r[pcol] == row[dcol]])
        for n in NAMES:
            cells_checked += 1
            if not eq(mval(row, n), e[n]):
                problems.append(f"{row[dcol]} / {n}: powerbi={mval(row, n)} expected={e[n]}")
    check(S, f"By {dcol}: {len(rows)} members x 33 measures == pandas", not problems, "; ".join(problems[:4]))
rows = qframe("by_dept_x_overtime")
problems = []
for row in rows:
    g = r[(r.Department == row["DimDepartment[Department]"]) & (r.OverTime == row["DimEmployee[OverTime]"])]
    e = expected(g)
    for n in NAMES:
        cells_checked += 1
        if not eq(mval(row, n), e[n]): problems.append(f"{row['DimDepartment[Department]']}/{row['DimEmployee[OverTime]']} {n}")
check(S, f"Two filters at once (Department x OverTime, {len(rows)} combinations) == pandas", not problems and len(rows) == 6, "; ".join(problems[:4]))
check(S, f"Total measure values compared with independent pandas: {cells_checked:,}", cells_checked > 50000)

# ------------------------------------------------------------------ E. edge cases
S = "E. Edge cases"
empty = qframe("empty_selection")[0]
bad = [(n, mval(empty, n)) for n in NAMES if not eq(mval(empty, n), COMPANY if n == "Company Attrition Rate" else None)]
check(S, "Empty selection: every measure BLANK (no error), Company baseline still 16.1%", not bad, str(bad))
kf = {row["FactEmployee[OverTimeFlag]"]: row for row in qframe("keepfilters_overtime_col")}
check(S, "KEEPFILTERS: on an OverTimeFlag axis, 'Attrition Rate - Overtime' is blank for flag 0 and 30.5% for flag 1",
      kf[0]["[AR_OT]"] is None and eq(kf[1]["[AR_OT]"], rate(r[r.ot == 1])) and kf[1]["[AR_NoOT]"] is None
      and eq(kf[0]["[AR_NoOT]"], rate(r[r.ot == 0])), json.dumps({k: {c: v for c, v in d.items()} for k, d in kf.items()}))
check(S, "Counts return 0 (not blank) when nobody matches: Overtime Employees = 0 on the flag-0 row", kf[0]["[OT_Emp]"] == 0)
kt = {row["FactEmployee[TrainingTimesLastYear]"]: row for row in qframe("keepfilters_training_col")}
ok = kt[0]["[NoTrain]"] == int((r.TrainingTimesLastYear == 0).sum()) and all(kt[v]["[NoTrain]"] == 0 for v in range(1, 7)) \
     and all(eq(kt[v]["[High]"], 1.0 if v >= 5 else 0.0) for v in range(0, 7))
check(S, "KEEPFILTERS: 'Employees with No Training' = 54 only on the 0-trainings row, 0 elsewhere; '% High Training' 100% only for 5-6", ok)
sl = qframe("slicer_sales")[0]
sales = r[r.Department == "Sales"]
check(S, "Slicer Department = Sales: rate 20.6%, company baseline unchanged 16.1%, gap +4.5 pts",
      eq(sl["[Rate]"], rate(sales)) and eq(sl["[Company]"], COMPANY) and eq(sl["[Vs]"], rate(sales) - COMPANY) and sl["[HC]"] == len(sales), str(sl))
nan = []
for k, q in live["queries"].items():
    for row in q.get("rows", []):
        for v in row:
            if isinstance(v, float) and (math.isnan(v) or math.isinf(v)): nan.append(k)
check(S, "No NaN or infinity in any query result", not nan, str(set(nan)))
check(S, "Negative control: a query using a non-existent column FAILS (proves errors are detected)",
      "error" in live["queries"]["negative_control_must_fail"], str(live["queries"]["negative_control_must_fail"])[:200])

# ------------------------------------------------------------------ F. report + UI
S = "F. Report + UI"
sres = json.load(open(SCH, encoding="utf-8-sig"))
check(S, f"All {len(sres)} report/project JSON files still pass Microsoft's official schemas", sres and all(x["ok"] for x in sres))
ui = json.load(open(UI, encoding="utf-8-sig"))
check(S, "Power BI window is open on the project", any("Talent_Learning_Dashboard" in w for w in ui["windows"]), str(ui["windows"]))
problems_counter = [t for t in ui["suspicious"] if re.fullmatch(r"Problems \d+", t)]
other = [t for t in ui["suspicious"] if t not in problems_counter]
check(S, "Power BI's own Problems counter shows 0 (or is not displayed)", all(t == "Problems 0" for t in problems_counter), str(problems_counter))
check(S, "No other dialog / warning / error text in the Power BI window", not other, str(other))
check(S, "No modal dialog open (only Power BI's built-in DAX-query / TMDL view panes)",
      all(d in ("Run DAX queries on your model", "Edit your model with TMDL") for d in ui["dialogs"]), str(ui["dialogs"]))
check(S, "Data pane lists the _Measures table", ui["measures_table_visible"])

passed, total_n = sum(res), len(res)
text = f"STEP 5 VALIDATION: {passed}/{total_n} checks passed\n\n" + "\n".join(out)
print(text)
open(REPORT, "w", encoding="utf-8").write(text + "\n")
sys.exit(0 if passed == total_n else 1)
