"""
Step 3 validation - Power Query layer.
  A. Syntax      : every .pq file parsed by Microsoft's official Power Query parser (+ negative control)
  B. Structure   : query names, dependency graph, column/type lists vs. source files and design
  C. Mirror      : Python re-implementation reconciles with Step 2 (validated) numbers
  D. Engine      : output of the REAL Power Query engine (Excel) == Python mirror, cell by cell
  E. QA_Checks   : in-model QA query ran in the engine, all PASS; expected values == mirror
  F. Model ready : keys unique, no nulls, referential integrity on the engine output
Usage: python validate_step3.py <project_root> <parse_results.json> <excel_engine_run.xlsx> <report.txt>
"""
import json, os, re, sys
import pandas as pd
from openpyxl import load_workbook

ROOT, PARSE_JSON, ENGINE_XLSX, REPORT = sys.argv[1:5]
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import mirror_step3 as mirror
mirror.CSV = os.path.join(ROOT, "01_Data/raw/WA_Fn-UseC_-HR-Employee-Attrition.csv")
mirror.LKP = os.path.join(ROOT, "01_Data/lookups/HR_Lookups.xlsx")
PQ_DIR = os.path.join(ROOT, "03_PowerBI/PowerQuery")
PROFILING = os.path.join(ROOT, "02_Excel_Profiling/HR_Data_Profiling.xlsx")

res, out = [], []
def check(sec, name, ok, detail=""):
    res.append(bool(ok)); out.append(f"[{'PASS' if ok else 'FAIL'}] {sec}: {name}" + (f" - {detail}" if detail and not ok else (f" - {detail}" if detail else "")))

LOADED = ["DimDepartment", "DimJobRole", "DimEducation", "DimJobSatisfaction", "DimAgeBand", "DimTenureBand",
          "DimIncomeBand", "DimTrainingBand", "DimEmployee", "FactEmployee", "QA_Checks"]
STAGING = ["fnGetLookup", "fnBandKey", "src_HR", "lkp_Education", "lkp_Satisfaction", "lkp_WorkLifeBalance",
           "lkp_PerformanceRating", "lkp_BusinessTravel", "lkp_AgeBand", "lkp_TenureBand", "lkp_IncomeBand",
           "lkp_TrainingBand", "stg_HR"]

# ------------------------------------------------------------------ A. syntax
S = "A. Syntax"
parsed = json.load(open(PARSE_JSON))
files = sorted(f for f in os.listdir(PQ_DIR) if f.endswith(".pq"))
check(S, f"{len(files)} query files found (expected 25)", len(files) == 25)
bad = [p["file"] + ": " + p["msg"] for p in parsed["main"] if not p["ok"]]
check(S, "Every query parses with @microsoft/powerquery-parser", not bad and len(parsed["main"]) == len(files), "; ".join(bad))
check(S, "Negative control: parser rejects deliberately broken M", parsed["negative"] and not parsed["negative"][0]["ok"])

# ------------------------------------------------------------------ B. structure
S = "B. Structure"
texts, names = {}, {}
for f in files:
    t = open(os.path.join(PQ_DIR, f), encoding="utf-8").read()
    q = re.sub(r"^\d+_", "", f[:-3])
    texts[q] = t
    first = t.splitlines()[0]
    names[q] = first
check(S, "Header comment of every file names its query", all(names[q].startswith(f"// {q} ") or names[q] == f"// {q}" for q in texts),
      str([q for q in texts if not names[q].startswith(f"// {q}")]))
check(S, "Query set = parameter + 13 staging + 11 loaded", set(texts) == set(["ProjectFolder"] + STAGING + LOADED),
      f"diff={set(texts) ^ set(['ProjectFolder'] + STAGING + LOADED)}")
roles_ok = all("(Load" in names[q] for q in LOADED) and all(("do not load" in names[q]) for q in STAGING)
check(S, "Header marks Load vs 'do not load' correctly", roles_ok)

def code_only(t):
    t = re.sub(r'"(?:[^"]|"")*"', '""', t)          # strip string literals
    return re.sub(r"//[^\n]*", "", t)               # strip comments
deps = {q: sorted({o for o in texts if o != q and re.search(rf"(?<![\w\[]){o}(?![\w\]])", code_only(t))}) for q, t in texts.items()}
def local_names(t):
    c = code_only(t)
    return set(re.findall(r"^\s*([A-Za-z_]\w*)\s*=", c, re.M)) | set(re.findall(r"[(,]\s*([A-Za-z_]\w*)\s+as\b", c))
undefined = {q: sorted({i for i in re.findall(r"(?<![\w\[])(?:lkp|stg|src|fn|Dim|Fact)[A-Za-z_]*\b(?![\]])", code_only(t))
                        if i not in texts and i not in local_names(t)}) for q, t in texts.items()}
undefined = {q: v for q, v in undefined.items() if v}
check(S, "Every referenced query exists", not undefined, str(undefined))
order = {re.sub(r"^\d+_", "", f[:-3]): i for i, f in enumerate(files)}
backward = [(q, d) for q, ds in deps.items() for d in ds if order[d] > order[q]]
check(S, "Dependencies only point to earlier files (no cycles; numbering = build order)", not backward, str(backward))
check(S, "Loaded tables never read files directly (only via staging)", all("File.Contents" not in texts[q] for q in LOADED))
check(S, "Only src_HR and fnGetLookup touch files", sorted(q for q in texts if "File.Contents" in texts[q]) == ["fnGetLookup", "src_HR"])
check(S, "Both file paths tolerate a missing trailing backslash", all('Text.EndsWith(ProjectFolder, "\\")' in texts[q] for q in ["fnGetLookup", "src_HR"]))

def type_pairs(t, step="Typed"):
    m = re.search(rf"{step}\s*=\s*Table\.TransformColumnTypes\([^,]+,\s*\{{(.*?)\}}\s*(?:,\s*\"en-US\")?\)", t, re.S)
    return re.findall(r'\{"([^"]+)",\s*(Int64\.Type|type text)\}', m.group(1)) if m else []

csv_cols = list(pd.read_csv(mirror.CSV, encoding="utf-8-sig", nrows=1).columns)
src_pairs = type_pairs(texts["src_HR"])
check(S, "src_HR types all 35 CSV columns, in CSV order", [c for c, _ in src_pairs] == csv_cols)
raw = pd.read_csv(mirror.CSV, encoding="utf-8-sig")
wrong = [c for c, ty in src_pairs if (ty == "Int64.Type") != pd.api.types.is_integer_dtype(raw[c])]
check(S, "src_HR type of each column matches the data (whole number vs text)", not wrong, str(wrong))
check(S, "src_HR uses UTF-8 (65001) so the BOM is removed", "Encoding = 65001" in texts["src_HR"])

lk = load_workbook(mirror.LKP, data_only=True)
lk_headers = {}
for ws in lk.worksheets:
    for tn in ws.tables:
        lk_headers[tn] = [c.value for c in ws[ws.tables[tn].ref][0]]
for q in [x for x in texts if x.startswith("lkp_")]:
    tn = re.search(r'fnGetLookup\("([^"]+)"\)', texts[q]).group(1)
    cols = [c for c, _ in type_pairs(texts[q])]
    check(S, f"{q}: table {tn} exists and typed columns = its headers", tn in lk_headers and cols == lk_headers[tn],
          f"{cols} vs {lk_headers.get(tn)}")

src, stg, T = mirror.build()
for q in LOADED[:-1]:
    pairs = type_pairs(texts[q])
    exp = [(c, "Int64.Type" if pd.api.types.is_integer_dtype(T[q][c]) else "type text") for c in T[q].columns]
    check(S, f"{q}: final column names + types match design", pairs == exp, f"M={pairs} design={exp}")
removed = re.search(r"Table\.RemoveColumns\(Source,\s*\{(.*?)\}", texts["stg_HR"], re.S).group(1)
check(S, "stg_HR drops exactly the 6 columns decided in Step 2",
      sorted(re.findall(r'"([^"]+)"', removed)) == sorted(["EmployeeCount", "Over18", "StandardHours", "DailyRate", "HourlyRate", "MonthlyRate"]))

# ------------------------------------------------------------------ C. mirror vs Step 2
S = "C. Mirror vs Step 2"
fe, de = T["FactEmployee"], T["DimEmployee"]
pv = load_workbook(PROFILING, data_only=True)
kpi = {pv["KPI_Summary"].cell(r, 1).value: pv["KPI_Summary"].cell(r, 2).value for r in range(5, 30) if pv["KPI_Summary"].cell(r, 1).value}
def close(a, b): return abs(float(a) - float(b)) < 1e-9
check(S, "Headcount = Step 2", len(fe) == kpi["Headcount"])
check(S, "Leavers = Step 2", fe.AttritionFlag.sum() == kpi["Leavers"])
check(S, "Attrition rate = Step 2", close(fe.AttritionFlag.mean(), kpi["Attrition Rate"]))
check(S, "Avg tenure = Step 2", close(fe.YearsAtCompany.mean(), kpi["Avg Tenure (years)"]))
check(S, "Training per employee = Step 2", close(fe.TrainingTimesLastYear.mean(), kpi["Training per Employee"]))
check(S, "Overtime % = Step 2", close(fe.OverTimeFlag.mean(), kpi["Overtime %"]))
check(S, "Avg job satisfaction = Step 2", close(fe.JobSatisfactionScore.mean(), kpi["Avg Job Satisfaction (1-4)"]))
check(S, "Avg monthly income = Step 2", close(fe.MonthlyIncome.mean(), kpi["Avg Monthly Income"]))
ap = pv["Attrition_Profile"]
prof = {}
sec = None
for r in range(1, ap.max_row + 1):
    a, d, e = ap.cell(r, 1).value, ap.cell(r, 4).value, ap.cell(r, 5).value
    is_note = bool(ap.cell(r, 1).font and ap.cell(r, 1).font.italic)
    if not is_note and (ap.cell(r + 1, 1).value == "Category" or (ap.cell(r + 2, 1).value == "Category" and ap.cell(r + 1, 1).font.italic)):
        sec = a
    if sec and isinstance(d, (int, float)) and a not in ("Total", "Category"):
        prof[(sec, a)] = (d, e)
j = fe.merge(T["DimDepartment"], on="DepartmentKey").merge(T["DimJobRole"], on="JobRoleKey") \
      .merge(T["DimTenureBand"], on="TenureBandKey").merge(T["DimAgeBand"], on="AgeBandKey").merge(T["DimIncomeBand"], on="IncomeBandKey")
for sec_name, col in [("Department", "Department"), ("Job Role", "JobRole"), ("Tenure Band (YearsAtCompany)", "TenureBand"),
                      ("Age Band", "AgeBand"), ("Monthly Income Band", "IncomeBand")]:
    g = j.groupby(col).AttritionFlag.agg(["count", "sum"])
    mism = [k for k in g.index if prof.get((sec_name, k)) != (g.loc[k, "count"], g.loc[k, "sum"])]
    check(S, f"Headcount + leavers by {col} via star-schema keys = Step 2 profile", not mism and len(g) > 0, str(mism))
tp = pv["Training_Profile"]
tb = {tp.cell(r, 1).value: (tp.cell(r, 4).value, tp.cell(r, 5).value) for r in range(1, 40) if tp.cell(r, 1).value in set(T["DimTrainingBand"].TrainingBand)}
g = fe.merge(T["DimTrainingBand"], on="TrainingBandKey").groupby("TrainingBand").AttritionFlag.agg(["count", "sum"])
check(S, "Headcount + leavers by TrainingBand = Step 2 profile", all(tb[k] == (g.loc[k, "count"], g.loc[k, "sum"]) for k in g.index))
check(S, "DimEmployee Status 'Left' count = leavers", (de.EmployeeStatus == "Left").sum() == kpi["Leavers"])
check(S, "DimEmployee OverTime 'Yes' = OverTimeFlag sum", (de.OverTime == "Yes").sum() == fe.OverTimeFlag.sum())
lab_ok = all(de[c].notna().all() for c in ["BusinessTravel", "EnvironmentSatisfaction", "RelationshipSatisfaction",
                                            "JobInvolvement", "WorkLifeBalance", "PerformanceRating"])
check(S, "All mapped labels in DimEmployee are filled", lab_ok)
emp_raw = raw.set_index("EmployeeNumber")
spot = all(emp_raw.loc[r.EmployeeID, "YearsAtCompany"] == r.YearsAtCompany and emp_raw.loc[r.EmployeeID, "Age"] == r.Age
           and (emp_raw.loc[r.EmployeeID, "Attrition"] == "Yes") == (r.AttritionFlag == 1) for r in fe.itertuples())
check(S, "Every fact row traces back to its raw CSV row (age, tenure, attrition)", spot)

# ------------------------------------------------------------------ D. engine vs mirror
S = "D. Engine vs Mirror"
eng = load_workbook(ENGINE_XLSX, data_only=True)
check(S, "Engine workbook has all 11 loaded queries as sheets", all(q[:31] in eng.sheetnames for q in LOADED),
      str([q for q in LOADED if q[:31] not in eng.sheetnames]))
E = {}
for q in LOADED:
    ws = eng[q[:31]]
    rows = list(ws.iter_rows(values_only=True))
    E[q] = pd.DataFrame(rows[1:], columns=rows[0])
for q in LOADED[:-1]:
    e, m = E[q], T[q]
    check(S, f"{q}: column names + order identical", list(e.columns) == list(m.columns), f"{list(e.columns)}")
    check(S, f"{q}: row count identical ({len(m)})", len(e) == len(m), f"engine={len(e)}")
    if list(e.columns) == list(m.columns) and len(e) == len(m):
        key = m.columns[0]
        es = e.sort_values(key).reset_index(drop=True); ms = m.sort_values(key).reset_index(drop=True)
        diff = 0
        for c in m.columns:
            a, b = es[c].tolist(), ms[c].tolist()
            diff += sum(1 for x, y in zip(a, b) if (str(int(x)) if isinstance(x, float) and float(x).is_integer() else str(x)) != str(y))
        check(S, f"{q}: all {m.size:,} cells identical", diff == 0, f"{diff} differing cells")
        nums = [c for c in m.columns if pd.api.types.is_integer_dtype(m[c])]
        bad_t = [c for c in nums if not all(isinstance(v, (int, float)) for v in e[c])]
        check(S, f"{q}: whole-number columns arrive as numbers (not text)", not bad_t, str(bad_t))

# ------------------------------------------------------------------ E. QA_Checks
S = "E. QA_Checks"
qa = E["QA_Checks"]
check(S, "QA_Checks returned 35 checks", len(qa) == 35, f"{len(qa)}")
fails = qa[qa.Result != "PASS"]
check(S, "All QA_Checks PASS inside the real engine", fails.empty, fails.to_string() if not fails.empty else "")
summ = {k: int(v) for k, v in {
    "Sum AttritionFlag (leavers)": fe.AttritionFlag.sum(), "Sum OverTimeFlag": fe.OverTimeFlag.sum(),
    "Sum TrainingTimesLastYear": fe.TrainingTimesLastYear.sum(), "Sum MonthlyIncome": fe.MonthlyIncome.sum(),
    "Sum YearsAtCompany": fe.YearsAtCompany.sum(), "Sum JobSatisfactionScore": fe.JobSatisfactionScore.sum(),
    "FactEmployee rows": len(fe), "DimEmployee rows": len(de), "FactEmployee columns": fe.shape[1], "DimEmployee columns": de.shape[1],
    **{f"{n} rows": len(T[n]) for n in LOADED[:8]}}.items()}
qa_exp = dict(zip(qa.Check, qa.Expected))
mism = {k: (qa_exp.get(k), v) for k, v in summ.items() if qa_exp.get(k) != v}
check(S, "Hardcoded expected values in QA_Checks = independent mirror", not mism, str(mism))

# ------------------------------------------------------------------ F. model readiness (engine output)
S = "F. Model readiness"
fact = E["FactEmployee"]
check(S, "FactEmployee grain: EmployeeID unique", fact.EmployeeID.is_unique)
check(S, "No nulls anywhere in the 10 model tables", all(E[q].notna().all().all() for q in LOADED[:-1]))
for dim, key in [("DimDepartment", "DepartmentKey"), ("DimJobRole", "JobRoleKey"), ("DimEducation", "EducationKey"),
                 ("DimJobSatisfaction", "JobSatisfactionKey"), ("DimAgeBand", "AgeBandKey"), ("DimTenureBand", "TenureBandKey"),
                 ("DimIncomeBand", "IncomeBandKey"), ("DimTrainingBand", "TrainingBandKey"), ("DimEmployee", "EmployeeID")]:
    d = E[dim]
    check(S, f"{dim}.{key} unique (can be the 1-side of a relationship)", d[key].is_unique)
    orphans = set(fact[key]) - set(d[key])
    check(S, f"FactEmployee.{key} -> {dim}: no orphan keys", not orphans, str(orphans))
for dim, lab, key in [("DimAgeBand", "AgeBand", "AgeBandKey"), ("DimTenureBand", "TenureBand", "TenureBandKey"),
                      ("DimIncomeBand", "IncomeBand", "IncomeBandKey"), ("DimTrainingBand", "TrainingBand", "TrainingBandKey")]:
    d = E[dim].sort_values(key)
    check(S, f"{dim}: sort key gives logical order", d.iloc[:, 2].is_monotonic_increasing)
de_e = E["DimEmployee"]
for lab in ["EnvironmentSatisfaction", "RelationshipSatisfaction", "JobInvolvement", "WorkLifeBalance", "PerformanceRating", "BusinessTravel"]:
    one = de_e.groupby(lab)[lab + "Sort"].nunique().max() == 1 and de_e.groupby(lab + "Sort")[lab].nunique().max() == 1
    check(S, f"DimEmployee.{lab} <-> {lab}Sort is 1:1 (Sort by column will work)", one)

passed, total = sum(res), len(res)
text = f"STEP 3 VALIDATION: {passed}/{total} checks passed\n\n" + "\n".join(out)
print(text)
open(REPORT, "w", encoding="utf-8").write(text + "\n")
sys.exit(0 if passed == total else 1)
