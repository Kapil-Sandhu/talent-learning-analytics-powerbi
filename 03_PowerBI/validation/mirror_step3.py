"""
Step 3 reference implementation (Python mirror of the Power Query logic).
Re-implements every query in 03_PowerBI/PowerQuery/*.pq with pandas, step by step, so the
expected output of each loaded table is known BEFORE it is loaded into Power BI.

Run from the project root:
    python 03_PowerBI/validation/mirror_step3.py
Writes expected tables to 03_PowerBI/validation/expected/*.csv and a schema/summary JSON.
"""
import json, os, sys
import pandas as pd

ROOT = sys.argv[1] if len(sys.argv) > 1 else "."
CSV = os.path.join(ROOT, "01_Data/raw/WA_Fn-UseC_-HR-Employee-Attrition.csv")
LKP = os.path.join(ROOT, "01_Data/lookups/HR_Lookups.xlsx")
OUT = os.path.join(ROOT, "03_PowerBI/validation/expected")


def lookup(name):
    """fnGetLookup: read one Excel table by name."""
    from openpyxl import load_workbook
    wb = load_workbook(LKP, data_only=True)
    for ws in wb.worksheets:
        if name in ws.tables:
            rows = [[c.value for c in r] for r in ws[ws.tables[name].ref]]
            return pd.DataFrame(rows[1:], columns=rows[0])
    raise KeyError(f"Lookup not found: {name}")


def band_key(value, bands):
    """fnBandKey: exactly one band must match."""
    m = bands[(bands.Min <= value) & (value <= bands.Max)]
    if len(m) != 1:
        raise ValueError(f"Band lookup failed: {value} matches {len(m)} bands")
    return int(m.BandKey.iloc[0])


def build():
    # src_HR
    src = pd.read_csv(CSV, encoding="utf-8-sig")
    assert src.shape == (1470, 35)

    # lkp_*
    lkp_education = lookup("Lkp_Education")
    lkp_sat = lookup("Lkp_Satisfaction")
    lkp_wlb = lookup("Lkp_WorkLifeBalance")
    lkp_perf = lookup("Lkp_PerformanceRating")
    lkp_trav = lookup("Lkp_BusinessTravel")
    def band_lkp(name, label, mn, mx):
        t = lookup(name).rename(columns={label: "Band", mn: "Min", mx: "Max", "SortOrder": "BandKey"})
        return t[["Band", "Min", "Max", "BandKey"]]
    lkp_age = band_lkp("Lkp_AgeBand", "AgeBand", "MinAge", "MaxAge")
    lkp_ten = band_lkp("Lkp_TenureBand", "TenureBand", "MinYears", "MaxYears")
    lkp_inc = band_lkp("Lkp_IncomeBand", "IncomeBand", "MinIncome", "MaxIncome")
    lkp_trn = band_lkp("Lkp_TrainingBand", "TrainingBand", "MinTrainings", "MaxTrainings")

    # stg_HR
    s = src.drop(columns=["EmployeeCount", "Over18", "StandardHours", "DailyRate", "HourlyRate", "MonthlyRate"])
    s = s.rename(columns={"EmployeeNumber": "EmployeeID"})
    for c in ["Attrition", "BusinessTravel", "Department", "EducationField", "Gender", "JobRole", "MaritalStatus", "OverTime"]:
        s[c] = s[c].str.strip()
    s["AttritionFlag"] = (s.Attrition == "Yes").astype(int)
    s["OverTimeFlag"] = (s.OverTime == "Yes").astype(int)
    s["EmployeeStatus"] = s.AttritionFlag.map({1: "Left", 0: "Active"})
    s["AgeBandKey"] = s.Age.apply(band_key, bands=lkp_age)
    s["TenureBandKey"] = s.YearsAtCompany.apply(band_key, bands=lkp_ten)
    s["IncomeBandKey"] = s.MonthlyIncome.apply(band_key, bands=lkp_inc)
    s["TrainingBandKey"] = s.TrainingTimesLastYear.apply(band_key, bands=lkp_trn)
    sat = dict(zip(lkp_sat.RatingCode, lkp_sat.RatingLabel))
    s["EnvironmentSatisfactionLabel"] = s.EnvironmentSatisfaction.map(sat)
    s["RelationshipSatisfactionLabel"] = s.RelationshipSatisfaction.map(sat)
    s["JobInvolvementLabel"] = s.JobInvolvement.map(sat)
    s["WorkLifeBalanceLabel"] = s.WorkLifeBalance.map(dict(zip(lkp_wlb.WLBCode, lkp_wlb.WLBLabel)))
    s["PerformanceRatingLabel"] = s.PerformanceRating.map(dict(zip(lkp_perf.PerformanceCode, lkp_perf.PerformanceLabel)))
    s["BusinessTravelLabel"] = s.BusinessTravel.map(dict(zip(lkp_trav.BusinessTravelCode, lkp_trav.BusinessTravelLabel)))
    s["BusinessTravelSort"] = s.BusinessTravel.map(dict(zip(lkp_trav.BusinessTravelCode, lkp_trav.SortOrder)))
    stg = s.sort_values("EmployeeID").reset_index(drop=True)

    T = {}
    for col, key in [("Department", "DepartmentKey"), ("JobRole", "JobRoleKey")]:
        d = pd.DataFrame({col: sorted(stg[col].unique())})
        d.insert(0, key, range(1, len(d) + 1))
        T["Dim" + col] = d
    T["DimEducation"] = lkp_education.rename(columns={"EducationCode": "EducationKey"})[["EducationKey", "EducationLevel"]]
    T["DimJobSatisfaction"] = lkp_sat.rename(columns={"RatingCode": "JobSatisfactionKey", "RatingLabel": "JobSatisfaction"})[["JobSatisfactionKey", "JobSatisfaction"]]
    for b, l, mn, mx in [("AgeBand", lkp_age, "MinAge", "MaxAge"), ("TenureBand", lkp_ten, "MinYears", "MaxYears"),
                         ("IncomeBand", lkp_inc, "MinIncome", "MaxIncome"), ("TrainingBand", lkp_trn, "MinTrainings", "MaxTrainings")]:
        T["Dim" + b] = l.rename(columns={"BandKey": b + "Key", "Band": b, "Min": mn, "Max": mx})[[b + "Key", b, mn, mx]]

    de = stg[["EmployeeID", "EmployeeStatus", "Gender", "MaritalStatus", "EducationField", "BusinessTravelLabel",
              "BusinessTravelSort", "OverTime", "JobLevel", "StockOptionLevel",
              "EnvironmentSatisfactionLabel", "EnvironmentSatisfaction", "RelationshipSatisfactionLabel",
              "RelationshipSatisfaction", "JobInvolvementLabel", "JobInvolvement", "WorkLifeBalanceLabel",
              "WorkLifeBalance", "PerformanceRatingLabel", "PerformanceRating"]]
    de = de.rename(columns={c: c + "Sort" for c in ["EnvironmentSatisfaction", "RelationshipSatisfaction",
                                                     "JobInvolvement", "WorkLifeBalance", "PerformanceRating"]})
    de = de.rename(columns={"BusinessTravelLabel": "BusinessTravel",
                            **{c + "Label": c for c in ["EnvironmentSatisfaction", "RelationshipSatisfaction",
                                                         "JobInvolvement", "WorkLifeBalance", "PerformanceRating"]}})
    T["DimEmployee"] = de

    f = stg.merge(T["DimDepartment"], on="Department", how="left").merge(T["DimJobRole"], on="JobRole", how="left")
    f["EducationKey"] = f.Education
    f["JobSatisfactionKey"] = f.JobSatisfaction
    f = f.rename(columns={c: c + "Score" for c in ["JobSatisfaction", "EnvironmentSatisfaction", "RelationshipSatisfaction",
                                                    "JobInvolvement", "WorkLifeBalance", "PerformanceRating"]})
    fact_cols = ["EmployeeID", "DepartmentKey", "JobRoleKey", "EducationKey", "JobSatisfactionKey",
                 "AgeBandKey", "TenureBandKey", "IncomeBandKey", "TrainingBandKey",
                 "AttritionFlag", "OverTimeFlag", "Age", "MonthlyIncome", "PercentSalaryHike",
                 "DistanceFromHome", "NumCompaniesWorked", "TotalWorkingYears", "YearsAtCompany",
                 "YearsInCurrentRole", "YearsSinceLastPromotion", "YearsWithCurrManager",
                 "TrainingTimesLastYear", "JobSatisfactionScore", "EnvironmentSatisfactionScore",
                 "RelationshipSatisfactionScore", "JobInvolvementScore", "WorkLifeBalanceScore",
                 "PerformanceRatingScore"]
    T["FactEmployee"] = f[fact_cols].sort_values("EmployeeID").reset_index(drop=True)
    return src, stg, T


if __name__ == "__main__":
    src, stg, T = build()
    os.makedirs(OUT, exist_ok=True)
    summary = {}
    for name, t in T.items():
        t.to_csv(os.path.join(OUT, f"{name}.csv"), index=False)
        summary[name] = {"rows": len(t), "columns": list(t.columns),
                         "types": {c: ("Int64" if pd.api.types.is_integer_dtype(t[c]) else "text") for c in t.columns}}
    fe = T["FactEmployee"]
    summary["_totals"] = {
        "Sum AttritionFlag": int(fe.AttritionFlag.sum()), "Sum OverTimeFlag": int(fe.OverTimeFlag.sum()),
        "Sum TrainingTimesLastYear": int(fe.TrainingTimesLastYear.sum()), "Sum MonthlyIncome": int(fe.MonthlyIncome.sum()),
        "Sum YearsAtCompany": int(fe.YearsAtCompany.sum()), "Sum EmployeeID": int(fe.EmployeeID.sum()),
        "Sum JobSatisfactionScore": int(fe.JobSatisfactionScore.sum()),
    }
    json.dump(summary, open(os.path.join(OUT, "_expected_summary.json"), "w"), indent=2)
    for k, v in summary.items():
        print(k, v if k == "_totals" else (v["rows"], len(v["columns"])))
