"""
Independent expected values for every measure, computed with pandas straight from the RAW CSV and the
lookup workbook (no model, no Power Query). Shared by the Step 6 validator.
"""
import os, math
import pandas as pd
from openpyxl import load_workbook

_ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))
raw = pd.read_csv(os.path.join(_ROOT, "01_Data/raw/WA_Fn-UseC_-HR-Employee-Attrition.csv"), encoding="utf-8-sig")
_lk = load_workbook(os.path.join(_ROOT, "01_Data/lookups/HR_Lookups.xlsx"), data_only=True)
def _t(name):
    for ws in _lk.worksheets:
        if name in ws.tables:
            rows = [[c.value for c in r] for r in ws[ws.tables[name].ref]]
            return pd.DataFrame(rows[1:], columns=rows[0])
def _band(series, t):
    b = _t(t)
    out = []
    for v in series:
        m = b[(b.iloc[:, 1] <= v) & (v <= b.iloc[:, 2])]
        assert len(m) == 1
        out.append(m.iloc[0, 0])
    return out
r = raw.copy()
r["att"] = (r.Attrition == "Yes").astype(int); r["ot"] = (r.OverTime == "Yes").astype(int)
sat = dict(_t("Lkp_Satisfaction").iloc[:, :2].values)
r["EducationLevel"] = r.Education.map(dict(_t("Lkp_Education").iloc[:, :2].values))
r["JobSatLabel"] = r.JobSatisfaction.map(sat); r["EnvLabel"] = r.EnvironmentSatisfaction.map(sat)
r["WLBLabel"] = r.WorkLifeBalance.map(dict(_t("Lkp_WorkLifeBalance").iloc[:, :2].values))
r["PerfLabel"] = r.PerformanceRating.map(dict(_t("Lkp_PerformanceRating").iloc[:, :2].values))
r["TravelLabel"] = r.BusinessTravel.map(dict(_t("Lkp_BusinessTravel").iloc[:, :2].values))
r["AgeBand"] = _band(r.Age, "Lkp_AgeBand"); r["TenureBand"] = _band(r.YearsAtCompany, "Lkp_TenureBand")
r["IncomeBand"] = _band(r.MonthlyIncome, "Lkp_IncomeBand"); r["TrainingBand"] = _band(r.TrainingTimesLastYear, "Lkp_TrainingBand")
r["Status"] = r.att.map({1: "Left", 0: "Active"})
RAW = r
COMPANY = r.att.mean()
COLMAP = {"DimDepartment.Department": "Department", "DimJobRole.JobRole": "JobRole", "DimTenureBand.TenureBand": "TenureBand",
          "DimEmployee.OverTime": "OverTime", "DimAgeBand.AgeBand": "AgeBand", "DimEmployee.JobLevel": "JobLevel",
          "DimEmployee.StockOptionLevel": "StockOptionLevel", "DimEmployee.MaritalStatus": "MaritalStatus",
          "DimEmployee.BusinessTravel": "TravelLabel", "DimIncomeBand.IncomeBand": "IncomeBand",
          "DimTrainingBand.TrainingBand": "TrainingBand", "DimJobSatisfaction.JobSatisfaction": "JobSatLabel",
          "DimEmployee.WorkLifeBalance": "WLBLabel", "DimEmployee.EnvironmentSatisfaction": "EnvLabel",
          "DimEmployee.PerformanceRating": "PerfLabel", "DimEmployee.EmployeeID": "EmployeeNumber",
          "DimEmployee.EmployeeStatus": "Status", "DimEmployee.Gender": "Gender", "DimEducation.EducationLevel": "EducationLevel"}

def _rate(g): return (g.att.sum() / len(g)) if len(g) else None
def _mean(s): return float(s.mean()) if len(s) else None
def _div(a, b): return None if a is None or b is None or b == 0 else a / b
def expected(g):
    hc = len(g)
    if hc == 0:
        return {k: (COMPANY if k == "Company Attrition Rate" else None) for k in NAMES}
    L = int(g.att.sum())
    e = {"Headcount": hc, "Leavers": L, "Active Employees": hc - L, "Attrition Rate": L / hc, "Retention Rate": (hc - L) / hc,
         "Company Attrition Rate": COMPANY, "Attrition vs Company (pts)": L / hc - COMPANY,
         "Avg Tenure (Years)": _mean(g.YearsAtCompany), "Median Tenure (Years)": float(g.YearsAtCompany.median()),
         "Avg Tenure of Leavers (Years)": _mean(g[g.att == 1].YearsAtCompany),
         "Early Tenure Attrition Rate (0-1 yrs)": _rate(g[g.YearsAtCompany <= 1]),
         "Avg Years Since Promotion": _mean(g.YearsSinceLastPromotion),
         "Total Trainings": int(g.TrainingTimesLastYear.sum()), "Training per Employee": g.TrainingTimesLastYear.sum() / hc,
         "Employees with No Training": int((g.TrainingTimesLastYear == 0).sum()),
         "% No Training": (g.TrainingTimesLastYear == 0).sum() / hc, "% High Training (5-6)": (g.TrainingTimesLastYear >= 5).sum() / hc,
         "Attrition Rate - No Training": _rate(g[g.TrainingTimesLastYear == 0]),
         "Overtime Employees": int(g.ot.sum()), "Overtime %": g.ot.sum() / hc,
         "Attrition Rate - Overtime": _rate(g[g.ot == 1]), "Attrition Rate - No Overtime": _rate(g[g.ot == 0]),
         "Avg Job Satisfaction": _mean(g.JobSatisfaction), "% Low Job Satisfaction": (g.JobSatisfaction == 1).sum() / hc,
         "Attrition Rate - Low Job Satisfaction": _rate(g[g.JobSatisfaction == 1]),
         "Avg Environment Satisfaction": _mean(g.EnvironmentSatisfaction), "Avg Work-Life Balance": _mean(g.WorkLifeBalance),
         "Avg Job Involvement": _mean(g.JobInvolvement), "Avg Performance Rating": _mean(g.PerformanceRating),
         "Avg Age": _mean(g.Age), "Avg Monthly Income": _mean(g.MonthlyIncome), "Avg Salary Hike (%)": _mean(g.PercentSalaryHike)}
    e["Overtime Attrition Ratio"] = _div(e["Attrition Rate - Overtime"], e["Attrition Rate - No Overtime"])
    return e
NAMES = ["Headcount", "Leavers", "Active Employees", "Attrition Rate", "Retention Rate", "Company Attrition Rate",
         "Attrition vs Company (pts)", "Avg Tenure (Years)", "Median Tenure (Years)", "Avg Tenure of Leavers (Years)",
         "Early Tenure Attrition Rate (0-1 yrs)", "Avg Years Since Promotion", "Total Trainings", "Training per Employee",
         "Employees with No Training", "% No Training", "% High Training (5-6)", "Attrition Rate - No Training",
         "Overtime Employees", "Overtime %", "Attrition Rate - Overtime", "Attrition Rate - No Overtime",
         "Overtime Attrition Ratio", "Avg Job Satisfaction", "% Low Job Satisfaction", "Attrition Rate - Low Job Satisfaction",
         "Avg Environment Satisfaction", "Avg Work-Life Balance", "Avg Job Involvement", "Avg Performance Rating",
         "Avg Age", "Avg Monthly Income", "Avg Salary Hike (%)"]
assert set(expected(r)) == set(NAMES)

def eq(a, b):
    if b is None: return a is None
    if a is None: return False
    if isinstance(a, float) and (math.isnan(a) or math.isinf(a)): return False
    return abs(float(a) - float(b)) <= 1e-9 * max(1.0, abs(float(b)))
